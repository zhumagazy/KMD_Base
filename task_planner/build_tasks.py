"""Генератор Excel-задачника «Задачи и календарь 2026–2030».

Запуск:  python task_planner/build_tasks.py
Результат: task_planner/Задачник_2026-2030.xlsx
"""
import datetime as dt
import os
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

# PLANNER_SMALL=1 — уменьшенная копия для быстрой проверки формул
SMALL = os.environ.get("PLANNER_SMALL") == "1"
OUT = Path(__file__).with_name("Задачник_2026-2030.xlsx") if not SMALL else Path(os.environ["PLANNER_OUT"])

YEARS = [2026, 2027, 2028, 2029, 2030]
MONTHS = ["Январь", "Февраль", "Март", "Апрель", "Май", "Июнь", "Июль",
          "Август", "Сентябрь", "Октябрь", "Ноябрь", "Декабрь"]
WEEKDAYS = ["Пн", "Вт", "Ср", "Чт", "Пт", "Сб", "Вс"]

S_HOME = "🏠 Главная"
S_CAL = "📅 Календарь"
S_TASK = "✅ Задачи"
S_PLAN = "🕘 План дня"
S_SUM = "📊 Итоги"
S_MON = "🗓 По месяцам"
S_REF = "📚 Справочник"

T_FIRST, T_LAST = 5, (40 if SMALL else 1004)
E_FIRST, E_LAST = 5, (40 if SMALL else 3004)

GROUPS = ["💼 Работа", "🏠 Дом и быт", "💪 Спорт и здоровье", "📚 Учёба", "👪 Семья и друзья",
          "🎨 Хобби", "💰 Финансы", "🧘 Отдых", "🚀 Проекты"]
CATS = ["📌 Задача", "📞 Звонок", "🤝 Встреча", "🛒 Покупка", "🔁 Привычка", "📝 Документы",
        "🎉 Мероприятие", "💡 Идея"]
PRIORITY = ["🔥 Высокий", "⭐ Средний", "💤 Низкий"]
T_STATUS = ["⏳ Не начата", "🔄 В работе", "✅ Выполнена", "❌ Отменена"]
E_STATUS = ["⏰ Запланировано", "✅ Сделано", "❌ Пропущено"]
T_DONE, T_CANCEL = T_STATUS[2], T_STATUS[3]
E_DONE, E_SKIP = E_STATUS[1], E_STATUS[2]

C = {
    "home": "4B3F9E", "cal": "3A86FF", "task": "E76F51", "plan": "F4A261", "sum": "8338EC",
    "mon": "2A9D8F", "ref": "495057",
    "input_fill": "FFFBEA", "grid": "D9DCE3", "text": "2B2D42", "muted": "8D99AE",
    "pos": "1B9E5A", "neg": "D62839", "done_fill": "EFEFEF", "today": "FFF3BF",
    "slate": "4F6272", "cal_bg": "F4F5F7", "event": "EAF3FF",
}
TAB_SHOW, TAB_TECH = "2E9E5B", "D62839"
FONT = "Arial"
PCT = '0%;-0%;"–"'
HOURS = '0.0" ч";-0.0" ч";"–"'
INT = '0;-0;"–"'

thin = Side(style="thin", color=C["grid"])
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
NAME_REFS = {}


def font(size=10, bold=False, color=None, italic=False, strike=False):
    return Font(name=FONT, size=size, bold=bold, color=color or C["text"], italic=italic, strike=strike)


def fill(hex_color):
    return PatternFill("solid", start_color=hex_color, end_color=hex_color)


def q(sheet):
    return f"'{sheet}'"


def title(ws, text, subtitle, color, width_cols):
    last = get_column_letter(width_cols)
    ws.merge_cells(f"A1:{last}1")
    ws.merge_cells(f"A2:{last}2")
    ws["A1"] = text
    ws["A1"].font = Font(name=FONT, size=18, bold=True, color="FFFFFF")
    ws["A1"].alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[1].height = 36
    ws["A2"] = subtitle
    ws["A2"].font = font(9, italic=True, color=C["muted"])
    ws["A2"].alignment = Alignment(vertical="center", indent=1, wrap_text=True)
    ws.row_dimensions[2].height = 30
    for col in range(1, width_cols + 1):
        ws.cell(1, col).fill = fill(color)
    ws.sheet_view.showGridLines = False


def header(ws, row, headers, color, start_col=1):
    for i, h in enumerate(headers):
        cell = ws.cell(row, start_col + i, h)
        cell.fill = fill(color)
        cell.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[row].height = 34


def widths(ws, values, start=1):
    for i, w in enumerate(values):
        ws.column_dimensions[get_column_letter(start + i)].width = w


def add_name(wb, name, sheet, ref):
    NAME_REFS[name] = f"{q(sheet)}!{ref}"
    wb.defined_names[name] = DefinedName(name, attr_text=f"{q(sheet)}!{ref}")


def add_list(ws, rng, source, prompt=None):
    dv = DataValidation(type="list", formula1=NAME_REFS.get(source, source), allow_blank=True)
    dv.error = "Выберите значение из списка (его можно дополнить в листе «📚 Справочник»)"
    dv.errorTitle = "Значение не из справочника"
    dv.showErrorMessage = True
    if prompt:
        dv.prompt = prompt
        dv.showInputMessage = True
    ws.add_data_validation(dv)
    dv.add(rng)


def add_date_dv(ws, rng):
    dv = DataValidation(type="date", operator="between", formula1="DATE(2025,1,1)",
                        formula2="DATE(2035,12,31)", allow_blank=True)
    dv.error = "Введите дату, например 05.10.2026"
    dv.errorTitle = "Неверная дата"
    dv.showErrorMessage = True
    dv.prompt = "Дата, например 05.10.2026 (Ctrl+; — сегодня)"
    dv.showInputMessage = True
    ws.add_data_validation(dv)
    dv.add(rng)


def add_time_dv(ws, rng):
    dv = DataValidation(type="time", operator="between", formula1="0", formula2="0.999988",
                        allow_blank=True)
    dv.error = "Введите время, например 09:30"
    dv.errorTitle = "Неверное время"
    dv.showErrorMessage = True
    dv.prompt = "Время, например 09:30"
    dv.showInputMessage = True
    ws.add_data_validation(dv)
    dv.add(rng)


def hhmm(ref):
    """Время → «09:30» без TEXT() (форматы TEXT зависят от языка Excel)."""
    return f'RIGHT("0"&HOUR({ref}),2)&":"&RIGHT("0"&MINUTE({ref}),2)'


def line_chart(title_text, y_title, height=8, width=22):
    ch = LineChart()
    ch.title = title_text
    ch.style = 2
    ch.height, ch.width = height, width
    ch.y_axis.title = y_title
    ch.y_axis.majorGridlines.spPr = GraphicalProperties(ln=None)
    ch.legend.position = "b"
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    return ch


def color_series(ch, colors, dashed=()):
    for i, s in enumerate(ch.series):
        s.graphicalProperties.line.solidFill = colors[i]
        s.graphicalProperties.line.width = 28000
        s.smooth = False
        if i in dashed:
            s.graphicalProperties.line.dashStyle = "dash"


wb = Workbook()
home = wb.active
home.title = S_HOME
cal = wb.create_sheet(S_CAL)
task = wb.create_sheet(S_TASK)
plan = wb.create_sheet(S_PLAN)
summ = wb.create_sheet(S_SUM)
mon = wb.create_sheet(S_MON)
ref = wb.create_sheet(S_REF)

# ---------------------------------------------------------------- СПРАВОЧНИК
title(ref, "📚 Справочник",
      "Источник всех выпадающих списков. Дописывайте свои группы занятий и категории в пустые жёлтые "
      "строки — они сразу появятся в списках и сводках.", C["ref"], 12)
widths(ref, [24, 3, 22, 3, 14, 3, 16, 3, 18, 3, 9, 13])
R0 = 4


def ref_list(col, head, values, rows, name, editable=False):
    h = ref.cell(R0, col, head)
    h.fill = fill(C["ref"])
    h.font = Font(name=FONT, bold=True, color="FFFFFF")
    h.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    h.border = BORDER
    for i in range(rows):
        cell = ref.cell(R0 + 1 + i, col, values[i] if i < len(values) else None)
        cell.border = BORDER
        cell.font = font(color="0000FF" if editable else C["text"])
        if editable:
            cell.fill = fill(C["input_fill"])
    L = get_column_letter(col)
    add_name(wb, name, S_REF, f"${L}${R0 + 1}:${L}${R0 + rows}")


ref.row_dimensions[R0].height = 34
N_GROUP = N_CAT = 20
ref_list(1, "🗂 Группа занятий", GROUPS, N_GROUP, "lst_groups", editable=True)
ref_list(3, "🏷 Категория", CATS, N_CAT, "lst_cats", editable=True)
ref_list(5, "🎯 Приоритет", PRIORITY, 3, "lst_prio")
ref_list(7, "🚦 Статус задачи", T_STATUS, 4, "lst_tstatus")
ref_list(9, "🚦 Статус дела", E_STATUS, 3, "lst_estatus")
ref_list(11, "📆 Год", YEARS, 5, "lst_years")
ref_list(12, "🗓 Месяц", MONTHS, 12, "lst_months")
ref.cell(3, 1, "Группа — сфера жизни, категория — тип дела").font = font(8, italic=True, color=C["muted"])
ref.freeze_panes = "A5"

# ---------------------------------------------------------------- ЗАДАЧИ
title(task, "✅ Задачи — список дел и отметка выполнения",
      "Одна строка = одна задача. Заполните название, группу, категорию, приоритет и срок. Выполнили — "
      "поставьте статус «✅ Выполнена» (строка зачеркнётся) и по желанию дату выполнения. "
      "Задачи со сроком сами появляются в календаре.", C["task"], 10)
header(task, 4, ["№", "✏️ Задача", "🗂 Группа", "🏷 Категория", "🎯 Приоритет", "📅 Срок",
                 "🚦 Статус", "🏁 Выполнена\n(дата)", "💬 Заметки", "⏰ Сроки",
                 "служ. подпись", "служ. ключ", "служ. просрочка"], C["task"])
widths(task, [5, 36, 20, 16, 13, 12, 15, 12, 28, 16, 8, 8, 8])
for col in ("K", "L", "M"):
    task.column_dimensions[col].hidden = True

T_EX = [
    ("Сдать квартальный отчёт", GROUPS[0], CATS[5], PRIORITY[0], "2026-10-01", T_STATUS[1], None),
    ("Позвонить маме", GROUPS[4], CATS[1], PRIORITY[1], "2026-10-04", T_DONE, "2026-10-04"),
    ("Оплатить коммуналку", GROUPS[6], CATS[0], PRIORITY[0], "2026-10-05", T_STATUS[0], None),
    ("Купить подарок другу", GROUPS[4], CATS[3], PRIORITY[1], "2026-10-07", T_STATUS[0], None),
    ("Записаться к стоматологу", GROUPS[2], CATS[1], PRIORITY[1], "2026-10-06", T_STATUS[0], None),
    ("Закончить модуль курса Python", GROUPS[3], CATS[0], PRIORITY[1], "2026-10-10", T_STATUS[1], None),
    ("Генеральная уборка", GROUPS[1], CATS[0], PRIORITY[2], "2026-10-11", T_STATUS[0], None),
    ("Подготовить презентацию проекта", GROUPS[8], CATS[5], PRIORITY[0], "2026-10-14", T_STATUS[0], None),
    ("Продлить страховку авто", GROUPS[6], CATS[5], PRIORITY[0], "2026-10-20", T_STATUS[0], None),
    ("Купить билеты в театр", GROUPS[7], CATS[6], PRIORITY[2], "2026-10-24", T_STATUS[0], None),
    ("Разобрать гардероб", GROUPS[1], CATS[0], PRIORITY[2], "2026-09-28", T_DONE, "2026-09-27"),
    ("Идеи для блога", GROUPS[5], CATS[7], PRIORITY[2], None, T_STATUS[0], None),
]
for r in range(T_FIRST, T_LAST + 1):
    ex = T_EX[r - T_FIRST] if r - T_FIRST < len(T_EX) else None
    if ex:
        name, grp, cat_, pr, d, st, dd = ex
        for c, v in ((2, name), (3, grp), (4, cat_), (5, pr), (7, st)):
            task.cell(r, c, v)
        if d:
            task.cell(r, 6, dt.date.fromisoformat(d))
        if dd:
            task.cell(r, 8, dt.date.fromisoformat(dd))
    task.cell(r, 1, f'=IF(B{r}="","",COUNTA($B${T_FIRST}:B{r}))')
    open_ = f'AND(G{r}<>"{T_DONE}",G{r}<>"{T_CANCEL}")'
    task.cell(r, 10, f'=IF(OR(B{r}="",F{r}="",NOT({open_})),"",IF(F{r}<TODAY(),"❗ просрочено "&(TODAY()-F{r})&" дн.",'
                     f'IF(F{r}=TODAY(),"🔥 сегодня",IF(F{r}-TODAY()<=3,"⏳ через "&(F{r}-TODAY())&" дн.",""))))')
    task.cell(r, 11, f'=IF(B{r}="","",IF(G{r}="{T_DONE}","✔ ",IF(G{r}="{T_CANCEL}","✖ ",'
                     f'IF(AND(F{r}<>"",F{r}<TODAY()),"❗ ","☐ ")))&B{r}&IF(E{r}="{PRIORITY[0]}"," 🔥",""))')
    task.cell(r, 12, f'=IF(OR(B{r}="",F{r}="",G{r}="{T_CANCEL}"),"",F{r}&"#"&'
                     f'COUNTIFS($F${T_FIRST}:F{r},F{r},$B${T_FIRST}:B{r},"?*",$G${T_FIRST}:G{r},"<>{T_CANCEL}"))')
    task.cell(r, 13, f'=IF(AND(B{r}<>"",F{r}<>"",F{r}<TODAY(),{open_}),MAX($M$4:M{r - 1})+1,"")')
    for c in range(1, 14):
        cell = task.cell(r, c)
        cell.border = BORDER
        cell.font = font(color="0000FF" if 2 <= c <= 9 else C["muted"])
        cell.alignment = Alignment(vertical="center", horizontal="center" if c in (1, 5, 6, 8) else None)
    task.cell(r, 6).number_format = "DD.MM.YYYY"
    task.cell(r, 8).number_format = "DD.MM.YYYY"
    task.cell(r, 2).fill = fill(C["input_fill"])
rng = lambda col: f"{col}{T_FIRST}:{col}{T_LAST}"  # noqa: E731
add_list(task, rng("C"), "lst_groups")
add_list(task, rng("D"), "lst_cats")
add_list(task, rng("E"), "lst_prio")
add_list(task, rng("G"), "lst_tstatus", prompt="Выполнили? Выберите «✅ Выполнена» — строка зачеркнётся.")
add_date_dv(task, rng("F"))
add_date_dv(task, rng("H"))
for nm, col in (("t_name", "B"), ("t_group", "C"), ("t_cat", "D"), ("t_prio", "E"), ("t_date", "F"),
                ("t_status", "G"), ("t_done", "H"), ("t_label", "K"), ("t_key", "L"), ("t_od", "M")):
    add_name(wb, nm, S_TASK, f"${col}${T_FIRST}:${col}${T_LAST}")
data = f"A{T_FIRST}:J{T_LAST}"
task.conditional_formatting.add(data, FormulaRule(
    formula=[f'$G{T_FIRST}="{T_DONE}"'], stopIfTrue=True,
    font=Font(strike=True, color="8D99AE"), fill=fill(C["done_fill"])))
task.conditional_formatting.add(data, FormulaRule(
    formula=[f'$G{T_FIRST}="{T_CANCEL}"'], stopIfTrue=True,
    font=Font(strike=True, italic=True, color="C0C0C0")))
task.conditional_formatting.add(data, FormulaRule(
    formula=[f'ISNUMBER($M{T_FIRST})'], fill=fill("FFE9E9")))
task.conditional_formatting.add(data, FormulaRule(
    formula=[f'AND($F{T_FIRST}=TODAY(),$B{T_FIRST}<>"")'], fill=fill(C["today"])))
task.conditional_formatting.add(f"J{T_FIRST}:J{T_LAST}", FormulaRule(
    formula=[f'LEFT($J{T_FIRST},1)="❗"'], font=Font(bold=True, color=C["neg"])))
task.conditional_formatting.add(f"G{T_FIRST}:G{T_LAST}", FormulaRule(
    formula=[f'$G{T_FIRST}="{T_STATUS[1]}"'], font=Font(bold=True, color=C["cal"])))
task.conditional_formatting.add(f"E{T_FIRST}:E{T_LAST}", FormulaRule(
    formula=[f'$E{T_FIRST}="{PRIORITY[0]}"'], font=Font(bold=True, color=C["neg"])))
task.freeze_panes = "C5"
task.auto_filter.ref = f"A4:J{T_LAST}"
task["G4"].comment = Comment("Отметка выполнения: «✅ Выполнена» зачёркивает задачу здесь и в календаре.", "Задачник")

# ---------------------------------------------------------------- ПЛАН ДНЯ
title(plan, "🕘 План дня — расписание по времени",
      "Одна строка = одно дело в расписании: дата, начало, окончание, что делаете. Длительность считается "
      "сама. Можно привязать дело к задаче из списка. Пересечения по времени подсвечиваются оранжевым.",
      C["plan"], 12)
header(plan, 4, ["📅 Дата", "День", "🕘 Начало", "🕔 Конец", "⏱ Часов", "✏️ Дело / занятие",
                 "🗂 Группа", "🏷 Категория", "🔗 Задача (из списка)", "🚦 Статус", "💬 Заметка",
                 "⚠️ Проверка", "служ. подпись", "служ. ключ", "служ. время"], C["plan"])
widths(plan, [12, 6, 9, 9, 8, 30, 20, 16, 30, 16, 22, 16, 8, 8, 8])
for col in ("M", "N", "O"):
    plan.column_dimensions[col].hidden = True

E_EX = [
    ("2026-10-04", "08:00", "09:00", "Утренняя пробежка", GROUPS[2], CATS[4], None, E_DONE),
    ("2026-10-04", "11:00", "11:30", "Звонок маме", GROUPS[4], CATS[1], "Позвонить маме", E_DONE),
    ("2026-10-04", "19:00", "21:00", "Курс Python: модуль 4", GROUPS[3], CATS[0], "Закончить модуль курса Python", E_STATUS[0]),
    ("2026-10-05", "09:00", "10:00", "Планёрка", GROUPS[0], CATS[2], None, E_STATUS[0]),
    ("2026-10-05", "10:00", "13:00", "Работа над отчётом", GROUPS[0], CATS[0], "Сдать квартальный отчёт", E_STATUS[0]),
    ("2026-10-05", "12:30", "13:30", "Обед с коллегой", GROUPS[4], CATS[2], None, E_STATUS[0]),
    ("2026-10-05", "18:30", "20:00", "Тренажёрный зал", GROUPS[2], CATS[4], None, E_STATUS[0]),
    ("2026-10-06", "08:00", "09:00", "Утренняя пробежка", GROUPS[2], CATS[4], None, E_STATUS[0]),
    ("2026-10-07", "18:00", "19:00", "Выбрать подарок", GROUPS[4], CATS[3], "Купить подарок другу", E_STATUS[0]),
    ("2026-10-10", "10:00", "14:00", "Поездка за город", GROUPS[7], CATS[6], None, E_STATUS[0]),
    ("2026-10-02", "19:00", "20:30", "Тренажёрный зал", GROUPS[2], CATS[4], None, E_DONE),
    ("2026-10-01", "15:00", "16:00", "Встреча с заказчиком", GROUPS[8], CATS[2], None, E_SKIP),
]


def t(s):
    h, m = map(int, s.split(":"))
    return dt.time(h, m)


for r in range(E_FIRST, E_LAST + 1):
    ex = E_EX[r - E_FIRST] if r - E_FIRST < len(E_EX) else None
    if ex:
        d, s1, s2, name, grp, cat_, tk, st = ex
        plan.cell(r, 1, dt.date.fromisoformat(d))
        plan.cell(r, 3, t(s1))
        plan.cell(r, 4, t(s2))
        for c, v in ((6, name), (7, grp), (8, cat_), (9, tk), (10, st)):
            plan.cell(r, c, v)
    plan.cell(r, 2, f'=IF(A{r}="","",CHOOSE(WEEKDAY(A{r},2),"Пн","Вт","Ср","Чт","Пт","Сб","Вс"))')
    plan.cell(r, 5, f'=IF(OR(C{r}="",D{r}=""),"",MOD(D{r}-C{r},1)*24)')
    plan.cell(r, 12, f'=IF(OR(A{r}="",C{r}="",D{r}=""),"",IF(COUNTIFS($A${E_FIRST}:$A${E_LAST},A{r},'
                     f'$C${E_FIRST}:$C${E_LAST},"<"&D{r},$D${E_FIRST}:$D${E_LAST},">"&C{r})>1,"⚠️ пересечение",""))')
    time_part = f'IF(C{r}="","",{hhmm(f"C{r}")}&IF(D{r}="","","–"&{hhmm(f"D{r}")})&" ")'
    plan.cell(r, 13, f'=IF(F{r}="","",IF(J{r}="{E_DONE}","✔ ",IF(J{r}="{E_SKIP}","✖ ","⏰ "))&{time_part}&F{r})')
    plan.cell(r, 15, f'=IF(A{r}="","",IF(C{r}="",0,C{r}))')
    plan.cell(r, 14, f'=IF(OR(A{r}="",F{r}=""),"",A{r}&"#"&(COUNTIFS($A${E_FIRST}:$A${E_LAST},A{r},'
                     f'$F${E_FIRST}:$F${E_LAST},"?*",$O${E_FIRST}:$O${E_LAST},"<"&O{r})'
                     f'+COUNTIFS($A${E_FIRST}:A{r},A{r},$F${E_FIRST}:F{r},"?*",$O${E_FIRST}:O{r},O{r})))')
    for c in range(1, 16):
        cell = plan.cell(r, c)
        cell.border = BORDER
        cell.font = font(color="0000FF" if c in (1, 3, 4, 6, 7, 8, 9, 10, 11) else C["muted"])
        cell.alignment = Alignment(vertical="center", horizontal="center" if c in (1, 2, 3, 4, 5) else None)
    plan.cell(r, 1).number_format = "DD.MM.YYYY"
    plan.cell(r, 3).number_format = "HH:MM"
    plan.cell(r, 4).number_format = "HH:MM"
    plan.cell(r, 5).number_format = '0.0;-0.0;"–"'
    for c in (1, 3, 4, 6):
        plan.cell(r, c).fill = fill(C["input_fill"])
rng = lambda col: f"{col}{E_FIRST}:{col}{E_LAST}"  # noqa: E731
add_date_dv(plan, rng("A"))
add_time_dv(plan, rng("C"))
add_time_dv(plan, rng("D"))
add_list(plan, rng("G"), "lst_groups")
add_list(plan, rng("H"), "lst_cats")
add_list(plan, rng("I"), "t_name", prompt="Необязательно: к какой задаче относится дело")
add_list(plan, rng("J"), "lst_estatus")
for nm, col in (("e_date", "A"), ("e_start", "C"), ("e_dur", "E"), ("e_name", "F"), ("e_group", "G"),
                ("e_cat", "H"), ("e_status", "J"), ("e_label", "M"), ("e_key", "N")):
    add_name(wb, nm, S_PLAN, f"${col}${E_FIRST}:${col}${E_LAST}")
data = f"A{E_FIRST}:L{E_LAST}"
plan.conditional_formatting.add(data, FormulaRule(
    formula=[f'$J{E_FIRST}="{E_DONE}"'], stopIfTrue=True,
    font=Font(strike=True, color="8D99AE"), fill=fill(C["done_fill"])))
plan.conditional_formatting.add(data, FormulaRule(
    formula=[f'$J{E_FIRST}="{E_SKIP}"'], stopIfTrue=True,
    font=Font(strike=True, italic=True, color="C0C0C0")))
plan.conditional_formatting.add(data, FormulaRule(
    formula=[f'AND($A{E_FIRST}=TODAY(),$F{E_FIRST}<>"")'], fill=fill(C["today"])))
plan.conditional_formatting.add(f"L{E_FIRST}:L{E_LAST}", FormulaRule(
    formula=[f'$L{E_FIRST}<>""'], font=Font(bold=True, color="E76F51"), fill=fill("FFF1EC")))
plan.conditional_formatting.add(f"B{E_FIRST}:B{E_LAST}", FormulaRule(
    formula=[f'OR($B{E_FIRST}="Сб",$B{E_FIRST}="Вс")'], font=Font(bold=True, color=C["neg"])))
plan.freeze_panes = "C5"
plan.auto_filter.ref = f"A4:L{E_LAST}"


# ---------------------------------------------------------------- общие формулы дня
def day_line(d, k, max_lines=None):
    """k-я строка списка дня d: сначала дела по времени, затем задачи. d — ссылка на ячейку с датой."""
    ne = f'COUNTIFS(e_date,{d},e_name,"?*")'
    nt = f'COUNTIFS(t_date,{d},t_name,"?*",t_status,"<>{T_CANCEL}")'
    item = (f'IF({k}<={ne},INDEX(e_label,MATCH({d}&"#"&{k},e_key,0)),'
            f'IF({k}-{ne}<={nt},INDEX(t_label,MATCH({d}&"#"&({k}-{ne}),t_key,0)),""))')
    if max_lines and k == max_lines:
        item = f'IF({ne}+{nt}>{max_lines},"… ещё "&({ne}+{nt}-{max_lines - 1}),{item})'
    return f'=IFERROR({item},"")'


def line_rules(ws, rng, first_cell):
    """Подсветка строк списка по значку в начале."""
    ws.conditional_formatting.add(rng, FormulaRule(
        formula=[f'LEFT({first_cell},1)="✔"'], font=Font(strike=True, color="8D99AE")))
    ws.conditional_formatting.add(rng, FormulaRule(
        formula=[f'LEFT({first_cell},1)="✖"'], font=Font(strike=True, italic=True, color="C0C0C0")))
    ws.conditional_formatting.add(rng, FormulaRule(
        formula=[f'LEFT({first_cell},1)="❗"'], font=Font(bold=True, color=C["neg"])))
    ws.conditional_formatting.add(rng, FormulaRule(
        formula=[f'LEFT({first_cell},1)="⏰"'], font=Font(color="1D4ED8")))
    ws.conditional_formatting.add(rng, FormulaRule(
        formula=[f'LEFT({first_cell},1)="…"'], font=Font(italic=True, bold=True, color=C["muted"])))


# ---------------------------------------------------------------- КАЛЕНДАРЬ
title(cal, "📅 Календарь",
      "Выберите год и месяц в жёлтых ячейках (пусто = текущий месяц). В каждом дне — дела из «🕘 План дня» "
      "(по времени) и задачи из «✅ Задачи» (по сроку). ✔ — сделано, ❗ — просрочено, ☐ — в планах.",
      C["cal"], 17)
widths(cal, [2] + [25] * 7 + [2, 24, 14])
CW = 7
cal["B3"] = "📆 Год ▸"
cal["D3"] = "🗓 Месяц ▸"
for a in ("B3", "D3"):
    cal[a].font = font(10, True, C["cal"])
    cal[a].alignment = Alignment(horizontal="right", vertical="center")
for a, src, hint in (("C3", "lst_years", "Год (пусто = текущий)"), ("E3", "lst_months", "Месяц (пусто = текущий)")):
    cal[a].fill = fill(C["input_fill"])
    cal[a].font = Font(name=FONT, size=11, bold=True, color="0000FF")
    cal[a].alignment = Alignment(horizontal="center", vertical="center")
    cal[a].border = Border(left=thin, right=thin, top=thin, bottom=Side(style="medium", color=C["cal"]))
    add_list(cal, a, src, prompt=hint)
cal.row_dimensions[3].height = 24
# служебные: выбранные год и месяц
cal["J3"] = '=IF($C$3="",YEAR(TODAY()),$C$3)'
cal["K3"] = '=IF($E$3="",MONTH(TODAY()),MATCH($E$3,lst_months,0))'
cal["J3"].font = font(8, color=C["muted"])
cal["K3"].font = font(8, color=C["muted"])
add_name(wb, "cal_y", S_CAL, "$J$3")
add_name(wb, "cal_m", S_CAL, "$K$3")
cal.merge_cells("B4:H4")
cal["B4"] = '=INDEX(lst_months,cal_m)&" "&cal_y'
cal["B4"].font = Font(name=FONT, size=20, bold=True, color=C["slate"])
cal["B4"].alignment = Alignment(horizontal="center", vertical="center")
cal.row_dimensions[4].height = 34
for j, wd in enumerate(WEEKDAYS):
    c = cal.cell(5, 2 + j, wd)
    c.font = Font(name=FONT, size=11, bold=True, color=C["neg"] if j >= 5 else C["text"])
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.fill = fill("E9ECF1")
    c.border = Border(bottom=Side(style="medium", color=C["slate"]))
cal.row_dimensions[5].height = 24

LINES = 6
BLOCK = LINES + 1
CAL_TOP = 6
START = "(DATE(cal_y,cal_m,1)-WEEKDAY(DATE(cal_y,cal_m,1),2)+1)"
medium = Side(style="medium", color="C5CBD3")
for w in range(6):
    r0 = CAL_TOP + w * BLOCK
    cal.row_dimensions[r0].height = 22
    for j in range(7):
        col = 2 + j
        L = get_column_letter(col)
        dcell = cal.cell(r0, col, f"={START}+{w * 7 + j}")
        dcell.number_format = "d"
        dcell.font = Font(name=FONT, size=13, bold=True, color=C["neg"] if j >= 5 else C["text"])
        dcell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        dcell.fill = fill(C["cal_bg"])
        dcell.border = Border(left=medium, right=medium, top=medium)
        dref = f"{L}${r0}"
        for k in range(1, LINES + 1):
            c = cal.cell(r0 + k, col, day_line(dref, k, LINES))
            c.font = font(8)
            c.alignment = Alignment(vertical="center", shrink_to_fit=True)
            c.fill = fill("FFFFFF")
            c.border = Border(left=medium, right=medium,
                              bottom=medium if k == LINES else None)
        # подсветка блока дня
        items = f"{L}{r0 + 1}:{L}{r0 + LINES}"
        cal.conditional_formatting.add(f"{L}{r0}", FormulaRule(
            formula=[f"{L}{r0}=TODAY()"], stopIfTrue=True,
            font=Font(bold=True, color="FFFFFF"), fill=fill(C["slate"])))
        cal.conditional_formatting.add(f"{L}{r0}:{L}{r0 + LINES}", FormulaRule(
            formula=[f"MONTH({dref})<>cal_m"], stopIfTrue=True,
            font=Font(color="C0C4CC"), fill=fill("FAFAFB")))
        cal.conditional_formatting.add(items, FormulaRule(
            formula=[f"{dref}=TODAY()"], fill=fill(C["today"])))
        line_rules(cal, items, f"{L}{r0 + 1}")
    for k in range(1, LINES + 1):
        cal.row_dimensions[r0 + k].height = 15
CAL_LAST = CAL_TOP + 6 * BLOCK - 1

# боковая панель: итоги месяца + легенда
side = [("📊 Этот месяц", None, None),
        ("Задач со сроком", f'=COUNTIFS(t_date,">="&DATE(cal_y,cal_m,1),t_date,"<"&DATE(cal_y,cal_m+1,1),t_name,"?*",t_status,"<>{T_CANCEL}")', INT),
        ("✅ Выполнено", f'=COUNTIFS(t_date,">="&DATE(cal_y,cal_m,1),t_date,"<"&DATE(cal_y,cal_m+1,1),t_status,"{T_DONE}")', INT),
        ("🎯 Прогресс", '=IF(K7=0,"",K8/K7)', PCT),
        ("❗ Просрочено", '=COUNTIFS(t_date,">="&DATE(cal_y,cal_m,1),t_date,"<"&DATE(cal_y,cal_m+1,1),t_od,">0")', INT),
        ("🕘 Дел в расписании", '=COUNTIFS(e_date,">="&DATE(cal_y,cal_m,1),e_date,"<"&DATE(cal_y,cal_m+1,1),e_name,"?*")', INT),
        ("⏱ Часов запланировано", '=SUMIFS(e_dur,e_date,">="&DATE(cal_y,cal_m,1),e_date,"<"&DATE(cal_y,cal_m+1,1))', HOURS),
        ("⏱ Часов сделано", f'=SUMIFS(e_dur,e_date,">="&DATE(cal_y,cal_m,1),e_date,"<"&DATE(cal_y,cal_m+1,1),e_status,"{E_DONE}")', HOURS)]
for i, (lbl, f, fmt) in enumerate(side):
    r = 6 + i
    a = cal.cell(r, 10, lbl)
    if f is None:
        a.font = Font(name=FONT, size=11, bold=True, color="FFFFFF")
        a.fill = fill(C["cal"])
        cal.cell(r, 11).fill = fill(C["cal"])
        continue
    a.font = font(9, True)
    a.border = BORDER
    v = cal.cell(r, 11, f)
    v.number_format = fmt
    v.font = Font(name=FONT, size=11, bold=True, color=C["cal"])
    v.alignment = Alignment(horizontal="center")
    v.border = BORDER
cal.conditional_formatting.add("K9", DataBarRule(start_type="num", start_value=0, end_type="num",
                                                 end_value=1, color="2E9E5B"))
cal.conditional_formatting.add("K10", CellIsRule(operator="greaterThan", formula=["0"],
                                                 font=Font(bold=True, color=C["neg"])))
leg = [("🔎 Обозначения", None), ("⏰ 09:00–10:00 — дело из расписания", "1D4ED8"),
       ("☐ задача в планах", C["text"]), ("❗ задача просрочена", C["neg"]),
       ("✔ сделано (зачёркнуто)", "8D99AE"), ("✖ пропущено / отменено", "C0C0C0"),
       ("🔥 высокий приоритет", C["text"]), ("… ещё N — не поместилось", C["muted"]),
       ("Тёмный кружок — сегодня", C["slate"])]
for i, (txt, color) in enumerate(leg):
    r = 16 + i
    c = cal.cell(r, 10, txt)
    if color is None:
        c.font = Font(name=FONT, size=11, bold=True, color="FFFFFF")
        c.fill = fill(C["cal"])
        cal.cell(r, 11).fill = fill(C["cal"])
    else:
        c.font = font(9, color=color, strike=txt.startswith(("✔", "✖")))
cal.cell(26, 10, "Больше 6 записей в день? Полный список — в «🕘 План дня» (фильтр по дате).").font = \
    font(8, italic=True, color=C["muted"])
cal.cell(26, 10).alignment = Alignment(wrap_text=True)
cal.merge_cells("J26:K28")
cal.freeze_panes = "A6"
cal.sheet_view.zoomScale = 90

# ---------------------------------------------------------------- ИТОГИ
title(summ, "📊 Итоги по задачам и делам",
      "Всё считается автоматически: статусы, группы занятий, категории и приоритеты. "
      "% выполнения считается без отменённых задач.", C["sum"], 9)
widths(summ, [24, 12, 12, 12, 12, 12, 12, 12, 3])


def table(r0, head_title, rows, key_name, color):
    """Таблица разреза: rows — формулы подписей; key_name — именованный диапазон признака у задач/дел."""
    summ.cell(r0 - 1, 1, head_title).font = font(13, True, color)
    header(summ, r0, ["", "Задач", "✅ Выполнено", "🎯 %", "❗ Просрочено", "🕘 Дел", "⏱ Часов\nплан",
                      "⏱ Часов\nсделано"], color)
    for i, lbl in enumerate(rows):
        r = r0 + 1 + i
        summ.cell(r, 1, lbl)
        e_key = {"t_group": "e_group", "t_cat": "e_cat", "t_prio": None}[key_name]
        summ.cell(r, 2, f'=IF($A{r}="","",COUNTIFS({key_name},$A{r},t_status,"<>{T_CANCEL}",t_name,"?*"))')
        summ.cell(r, 3, f'=IF($A{r}="","",COUNTIFS({key_name},$A{r},t_status,"{T_DONE}"))')
        summ.cell(r, 4, f'=IF(OR($A{r}="",N(B{r})=0),"",C{r}/B{r})')
        summ.cell(r, 5, f'=IF($A{r}="","",COUNTIFS({key_name},$A{r},t_od,">0"))')
        if e_key:
            summ.cell(r, 6, f'=IF($A{r}="","",COUNTIFS({e_key},$A{r},e_name,"?*"))')
            summ.cell(r, 7, f'=IF($A{r}="","",SUMIFS(e_dur,{e_key},$A{r}))')
            summ.cell(r, 8, f'=IF($A{r}="","",SUMIFS(e_dur,{e_key},$A{r},e_status,"{E_DONE}"))')
        for c in range(1, 9):
            cell = summ.cell(r, c)
            cell.border = BORDER
            cell.font = font(10, c == 1)
            cell.alignment = Alignment(horizontal="left" if c == 1 else "center")
        for c in (2, 3, 5, 6):
            summ.cell(r, c).number_format = INT
        summ.cell(r, 4).number_format = PCT
        summ.cell(r, 7).number_format = HOURS
        summ.cell(r, 8).number_format = HOURS
    last = r0 + len(rows)
    tr = last + 1
    summ.cell(tr, 1, "Σ Итого")
    for c in (2, 3, 5, 6, 7, 8):
        L = get_column_letter(c)
        summ.cell(tr, c, f"=SUM({L}{r0 + 1}:{L}{last})")
    summ.cell(tr, 4, f'=IF(B{tr}=0,"",C{tr}/B{tr})')
    for c in range(1, 9):
        cell = summ.cell(tr, c)
        cell.border = BORDER
        cell.font = font(10, True)
        cell.fill = fill("EDE3FF")
        cell.alignment = Alignment(horizontal="left" if c == 1 else "center")
    for c in (2, 3, 5, 6):
        summ.cell(tr, c).number_format = INT
    summ.cell(tr, 4).number_format = PCT
    summ.cell(tr, 7).number_format = HOURS
    summ.cell(tr, 8).number_format = HOURS
    summ.conditional_formatting.add(f"D{r0 + 1}:D{tr}", DataBarRule(
        start_type="num", start_value=0, end_type="num", end_value=1, color="2E9E5B"))
    summ.conditional_formatting.add(f"E{r0 + 1}:E{tr}", CellIsRule(
        operator="greaterThan", formula=["0"], font=Font(bold=True, color=C["neg"]), fill=fill("FFE9E9")))
    summ.conditional_formatting.add(f"G{r0 + 1}:G{last}", DataBarRule(
        start_type="num", start_value=0, end_type="max", color="F4A261"))
    return r0 + 1, last, tr


# Статусы — компактный блок
summ.cell(3, 1, "🚦 По статусам").font = font(13, True, C["sum"])
header(summ, 4, ["Статус", "Задач", "Доля"], C["sum"])
for i, st in enumerate(T_STATUS):
    r = 5 + i
    summ.cell(r, 1, st)
    summ.cell(r, 2, f'=COUNTIFS(t_status,$A{r},t_name,"?*")')
    summ.cell(r, 3, f'=IF(B$9=0,"",B{r}/B$9)')
    for c in range(1, 4):
        summ.cell(r, c).border = BORDER
        summ.cell(r, c).font = font(10, c == 1)
        summ.cell(r, c).alignment = Alignment(horizontal="left" if c == 1 else "center")
    summ.cell(r, 2).number_format = INT
    summ.cell(r, 3).number_format = PCT
summ.cell(9, 1, "Σ Всего задач")
summ.cell(9, 2, "=SUM(B5:B8)")
summ.cell(10, 1, "❗ Из них просрочено")
summ.cell(10, 2, '=COUNTIFS(t_od,">0")')
summ.cell(11, 1, "📭 Без срока")
summ.cell(11, 2, '=COUNTIFS(t_name,"?*")-COUNTIFS(t_name,"?*",t_date,">0")')
for r in (9, 10, 11):
    for c in (1, 2):
        summ.cell(r, c).border = BORDER
        summ.cell(r, c).font = font(10, True)
        summ.cell(r, c).alignment = Alignment(horizontal="left" if c == 1 else "center")
    summ.cell(r, 2).number_format = INT
summ.conditional_formatting.add("C5:C8", DataBarRule(start_type="num", start_value=0, end_type="num",
                                                     end_value=1, color="8338EC"))
summ.conditional_formatting.add("B10", CellIsRule(operator="greaterThan", formula=["0"],
                                                  font=Font(bold=True, color=C["neg"])))

g_rows = [f'=IF({q(S_REF)}!A{R0 + 1 + i}="","",{q(S_REF)}!A{R0 + 1 + i})' for i in range(12)]
c_rows = [f'=IF({q(S_REF)}!C{R0 + 1 + i}="","",{q(S_REF)}!C{R0 + 1 + i})' for i in range(12)]
G_FIRST, G_LAST, G_TOT = table(15, "🗂 По группам занятий", g_rows, "t_group", C["sum"])
C_FIRST, C_LAST, C_TOT = table(G_TOT + 4, "🏷 По категориям", c_rows, "t_cat", "2A9D8F")
P_FIRST, P_LAST, P_TOT = table(C_TOT + 4, "🎯 По приоритетам", PRIORITY, "t_prio", "E76F51")
summ.freeze_panes = "B3"

bc = BarChart()
bc.type = "bar"
bc.style = 2
bc.title = "🗂 Задачи по группам: всего и выполнено"
bc.height, bc.width = 9, 16
bc.add_data(Reference(summ, min_col=2, max_col=3, min_row=15, max_row=G_LAST), titles_from_data=True)
bc.set_categories(Reference(summ, min_col=1, min_row=G_FIRST, max_row=G_LAST))
bc.series[0].graphicalProperties.solidFill = "C9B6F7"
bc.series[1].graphicalProperties.solidFill = "2E9E5B"
bc.legend.position = "b"
bc.y_axis.majorGridlines.spPr = GraphicalProperties(ln=None)
bc.x_axis.delete = False
bc.y_axis.delete = False
summ.add_chart(bc, "J3")
bc2 = BarChart()
bc2.type = "bar"
bc2.style = 2
bc2.title = "⏱ Часы по группам: план и сделано"
bc2.height, bc2.width = 9, 16
bc2.add_data(Reference(summ, min_col=7, max_col=8, min_row=15, max_row=G_LAST), titles_from_data=True)
bc2.set_categories(Reference(summ, min_col=1, min_row=G_FIRST, max_row=G_LAST))
bc2.series[0].graphicalProperties.solidFill = "F4C9A0"
bc2.series[1].graphicalProperties.solidFill = "E76F51"
bc2.legend.position = "b"
bc2.y_axis.majorGridlines.spPr = GraphicalProperties(ln=None)
bc2.x_axis.delete = False
bc2.y_axis.delete = False
summ.add_chart(bc2, "J22")

# ---------------------------------------------------------------- ПО МЕСЯЦАМ
title(mon, "🗓 По месяцам: январь 2026 → декабрь 2030",
      "Задачи считаются по месяцу срока, «Закрыто» — по дате выполнения. Дела и часы — из «🕘 План дня». "
      "Текущий месяц подсвечен жёлтым.", C["mon"], 12)
header(mon, 4, ["📆 Год", "🗓 Месяц", "начало", "Подпись", "📌 Задач\nсо сроком", "✅ Выполнено",
                "🎯 %", "❗ Просрочено", "🏁 Закрыто\nв месяце", "🕘 Дел", "⏱ Часов\nплан", "⏱ Часов\nсделано"],
       C["mon"])
widths(mon, [8, 12, 11, 9, 12, 12, 9, 12, 12, 9, 11, 11])
mon.column_dimensions["C"].hidden = True
mon.column_dimensions["D"].hidden = True
M_FIRST = 5
r = M_FIRST
for y in YEARS:
    for mi, m in enumerate(MONTHS, start=1):
        mon.cell(r, 1, y)
        mon.cell(r, 2, m)
        mon.cell(r, 3, dt.date(y, mi, 1))
        mon.cell(r, 4, f'=LEFT(B{r},3)&" "&RIGHT(A{r},2)')
        rngd = f'">="&C{r},t_date,"<"&EDATE(C{r},1)'
        mon.cell(r, 5, f'=COUNTIFS(t_date,{rngd},t_name,"?*",t_status,"<>{T_CANCEL}")')
        mon.cell(r, 6, f'=COUNTIFS(t_date,{rngd},t_status,"{T_DONE}")')
        mon.cell(r, 7, f'=IF(E{r}=0,"",F{r}/E{r})')
        mon.cell(r, 8, f'=COUNTIFS(t_date,{rngd},t_od,">0")')
        mon.cell(r, 9, f'=COUNTIFS(t_done,">="&C{r},t_done,"<"&EDATE(C{r},1))')
        erng = f'e_date,">="&C{r},e_date,"<"&EDATE(C{r},1)'
        mon.cell(r, 10, f'=COUNTIFS({erng},e_name,"?*")')
        mon.cell(r, 11, f'=SUMIFS(e_dur,{erng})')
        mon.cell(r, 12, f'=SUMIFS(e_dur,{erng},e_status,"{E_DONE}")')
        for c in range(1, 13):
            cell = mon.cell(r, c)
            cell.border = BORDER
            cell.font = font()
            cell.alignment = Alignment(horizontal="center")
        for c in (5, 6, 8, 9, 10):
            mon.cell(r, c).number_format = INT
        mon.cell(r, 7).number_format = PCT
        mon.cell(r, 11).number_format = HOURS
        mon.cell(r, 12).number_format = HOURS
        if mi == 1 and y != YEARS[0]:
            for c in range(1, 13):
                mon.cell(r, c).border = Border(left=thin, right=thin, bottom=thin,
                                               top=Side(style="medium", color=C["mon"]))
        r += 1
M_LAST = r - 1
tr = M_LAST + 1
mon.cell(tr, 1, "Σ ИТОГО")
mon.merge_cells(start_row=tr, start_column=1, end_row=tr, end_column=4)
for c in (5, 6, 8, 9, 10, 11, 12):
    L = get_column_letter(c)
    mon.cell(tr, c, f"=SUM({L}{M_FIRST}:{L}{M_LAST})")
mon.cell(tr, 7, f'=IF(E{tr}=0,"",F{tr}/E{tr})')
for c in range(1, 13):
    cell = mon.cell(tr, c)
    cell.border = BORDER
    cell.font = font(10, True)
    cell.fill = fill("DDF2EF")
    cell.alignment = Alignment(horizontal="center")
for c in (5, 6, 8, 9, 10):
    mon.cell(tr, c).number_format = INT
mon.cell(tr, 7).number_format = PCT
mon.cell(tr, 11).number_format = HOURS
mon.cell(tr, 12).number_format = HOURS
mon.conditional_formatting.add(f"A{M_FIRST}:L{M_LAST}", FormulaRule(
    formula=[f"AND($A{M_FIRST}=YEAR(TODAY()),MONTH($C{M_FIRST})=MONTH(TODAY()))"],
    fill=fill(C["today"]), font=Font(bold=True)))
mon.conditional_formatting.add(f"G{M_FIRST}:G{tr}", DataBarRule(
    start_type="num", start_value=0, end_type="num", end_value=1, color="2E9E5B"))
mon.conditional_formatting.add(f"H{M_FIRST}:H{tr}", CellIsRule(
    operator="greaterThan", formula=["0"], font=Font(bold=True, color=C["neg"])))
mon.freeze_panes = "E5"
mon.auto_filter.ref = f"A4:L{M_LAST}"
cats_ref = Reference(mon, min_col=4, min_row=M_FIRST, max_row=M_LAST)
ch = line_chart("📌 Задачи по месяцам: запланировано vs выполнено", "задач", height=9, width=26)
ch.add_data(Reference(mon, min_col=5, max_col=6, min_row=4, max_row=M_LAST), titles_from_data=True)
ch.set_categories(cats_ref)
color_series(ch, ["8338EC", "2E9E5B"], dashed=[0])
mon.add_chart(ch, "N4")
ch = line_chart("⏱ Часы по месяцам: план vs сделано", "часов", height=9, width=26)
ch.add_data(Reference(mon, min_col=11, max_col=12, min_row=4, max_row=M_LAST), titles_from_data=True)
ch.set_categories(cats_ref)
color_series(ch, ["F4A261", "E76F51"], dashed=[0])
mon.add_chart(ch, "N23")

# ---------------------------------------------------------------- ГЛАВНАЯ
title(home, "🏠 Задачи и календарь 2026 – 2030",
      "Личный задачник: задачи с отметкой выполнения, расписание по времени, календарь месяца и итоги.",
      C["home"], 11)
widths(home, [3, 30, 14, 3, 40, 3, 30, 14, 3, 3, 3])


def block(row, col, text, color):
    c = home.cell(row, col, text)
    c.font = Font(name=FONT, size=12, bold=True, color="FFFFFF")
    c.fill = fill(color)
    home.cell(row, col + 1).fill = fill(color)


block(4, 2, "📌 Сводка", C["home"])
kpi = [("Всего задач", f"={q(S_SUM)}!B9", INT),
       ("✅ Выполнено", f"={q(S_SUM)}!B7", INT),
       ("🎯 Прогресс", f'=IF(({q(S_SUM)}!B9-{q(S_SUM)}!B8)=0,"",{q(S_SUM)}!B7/({q(S_SUM)}!B9-{q(S_SUM)}!B8))', PCT),
       ("❗ Просрочено", f"={q(S_SUM)}!B10", INT),
       ("🔥 На сегодня (дел + задач)", f'=COUNTIFS(e_date,TODAY(),e_name,"?*")+COUNTIFS(t_date,TODAY(),t_name,"?*",t_status,"<>{T_CANCEL}")', INT),
       ("⏱ Часов на этой неделе", '=SUMIFS(e_dur,e_date,">="&(TODAY()-WEEKDAY(TODAY(),2)+1),e_date,"<"&(TODAY()-WEEKDAY(TODAY(),2)+8))', HOURS)]
for i, (lbl, f, fmt) in enumerate(kpi):
    r = 5 + i
    home.cell(r, 2, lbl).font = font(10, True)
    v = home.cell(r, 3, f)
    v.number_format = fmt
    v.font = Font(name=FONT, size=13, bold=True, color=C["home"])
    v.alignment = Alignment(horizontal="center")
    for c in (2, 3):
        home.cell(r, c).border = BORDER
    home.row_dimensions[r].height = 22
home.conditional_formatting.add("C7", DataBarRule(start_type="num", start_value=0, end_type="num",
                                                  end_value=1, color="2E9E5B"))
home.conditional_formatting.add("C8", CellIsRule(operator="greaterThan", formula=["0"],
                                                 font=Font(bold=True, color=C["neg"]), fill=fill("FFE9E9")))

# Сегодня
block(4, 5, "", C["cal"])
home.cell(4, 5, '="☀️ Сегодня, "&RIGHT("0"&DAY(TODAY()),2)&"."&RIGHT("0"&MONTH(TODAY()),2)&"."&YEAR(TODAY())&" ("&CHOOSE(WEEKDAY(TODAY(),2),"Пн","Вт","Ср","Чт","Пт","Сб","Вс")&")"')
home.cell(4, 6).fill = fill("FFFFFF")
AG = 12
home.cell(5, 6, "=TODAY()").number_format = "DD.MM.YYYY"
home.cell(5, 6).font = font(8, color="FFFFFF")
for k in range(1, AG + 1):
    r = 4 + k
    c = home.cell(r, 5, day_line("$F$5", k, AG))
    c.font = font(10)
    c.border = Border(bottom=Side(style="hair", color=C["grid"]))
line_rules(home, f"E5:E{4 + AG}", "E5")
home.cell(5 + AG, 5, "Пусто? Добавьте дела в «🕘 План дня» и задачи со сроком на сегодня.").font = \
    font(8, italic=True, color=C["muted"])

# Просроченные
block(4, 7, "❗ Просроченные задачи", C["neg"])
OD = 10
for k in range(1, OD + 1):
    r = 4 + k
    m = f"MATCH({k},t_od,0)"
    home.cell(r, 7, f'=IFERROR(INDEX(t_name,{m}),"")').font = font(10, color=C["neg"])
    c = home.cell(r, 8, f'=IFERROR(INDEX(t_date,{m}),"")')
    c.number_format = "DD.MM.YYYY"
    c.font = font(9, color=C["muted"])
    c.alignment = Alignment(horizontal="center")
    for cc in (7, 8):
        home.cell(r, cc).border = Border(bottom=Side(style="hair", color=C["grid"]))
home.cell(5 + OD, 7, f'=IF(COUNT(t_od)>{OD},"… и ещё "&(COUNT(t_od)-{OD}),IF(COUNT(t_od)=0,"🎉 Просроченных нет!",""))').font = \
    font(9, True, C["muted"])

# Навигация
block(12, 2, "🧭 Навигация", C["home"])
nav = [(S_CAL, "календарь месяца"), (S_TASK, "задачи и отметка выполнения"), (S_PLAN, "расписание по времени"),
       (S_SUM, "итоги по группам и категориям"), (S_MON, "итоги по месяцам + графики"), (S_REF, "справочник списков")]
for i, (sh, desc) in enumerate(nav):
    c = home.cell(13 + i, 2, sh)
    c.hyperlink = f"#'{sh}'!A1"
    c.font = Font(name=FONT, size=11, bold=True, color="1155CC", underline="single")
    home.cell(13 + i, 3, desc).font = font(8, italic=True, color=C["muted"])

# Легенда
block(20, 2, "🎨 Как пользоваться", C["home"])
legend = [
    "1️⃣ «✅ Задачи» — записываете задачи: группа, категория, приоритет, срок. Выполнили → статус «✅ Выполнена».",
    "2️⃣ «🕘 План дня» — расписание: дата, начало, конец, дело. Сделали → статус «✅ Сделано».",
    "3️⃣ «📅 Календарь» — выберите год и месяц: в каждом дне видны дела и задачи.",
    "4️⃣ «📊 Итоги» и «🗓 По месяцам» — считаются сами, вводить ничего не нужно.",
    "🟢 Зелёные вкладки — витрины (смотреть), 🔴 красные — технические (вводить данные).",
    "Жёлтые ячейки — ввод. Синий шрифт — ваши данные. Зачёркнуто — выполнено. Красным — просрочено.",
    "В «Задачах» и «Плане дня» есть строки-ПРИМЕРЫ — замените или удалите их.",
]
for i, txt in enumerate(legend):
    r = 21 + i
    home.merge_cells(start_row=r, start_column=2, end_row=r, end_column=8)
    c = home.cell(r, 2, txt)
    c.font = font(10)
    c.alignment = Alignment(vertical="center", wrap_text=True)
    home.row_dimensions[r].height = 18
ch = line_chart("📌 Задачи по месяцам: запланировано vs выполнено", "задач", height=8, width=24)
ch.add_data(Reference(mon, min_col=5, max_col=6, min_row=4, max_row=M_LAST), titles_from_data=True)
ch.set_categories(cats_ref)
color_series(ch, ["8338EC", "2E9E5B"], dashed=[0])
home.add_chart(ch, "B30")

# ---------------------------------------------------------------- вкладки
for ws_ in (home, cal, summ, mon):
    ws_.sheet_properties.tabColor = TAB_SHOW
for ws_ in (task, plan, ref):
    ws_.sheet_properties.tabColor = TAB_TECH

wb.active = 0
wb.save(OUT)
print(f"Saved {OUT}")
