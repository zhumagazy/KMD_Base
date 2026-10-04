"""Генератор Excel-планера «Финансовый план 2026–2030».

Запуск:  python finance_planner/build_planner.py
Результат: finance_planner/Финансовый_план_2026-2030.xlsx
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).with_name("Финансовый_план_2026-2030.xlsx")

YEARS = [2026, 2027, 2028, 2029, 2030]
MONTHS = ["Январь", "Февраль", "Март", "Апрель", "Май", "Июнь", "Июль",
          "Август", "Сентябрь", "Октябрь", "Ноябрь", "Декабрь"]

# Названия листов
S_HOME = "🏠 Главная"
S_EXP = "💸 Расходы"
S_INC = "💰 Доходы"
S_MON = "📅 По месяцам"
S_YEAR = "📊 По годам"
S_WISH = "✨ Хотелки"
S_REF = "📚 Справочник"

EXP_FIRST, EXP_LAST = 5, 504      # строки ввода расходов
INC_FIRST, INC_LAST = 5, 304      # строки ввода доходов
WISH_ROWS = 200

# Палитра
C = {
    "home": "4B3F9E", "exp": "E76F51", "inc": "2A9D8F", "mon": "3A86FF",
    "year": "8338EC", "wish": "FF5D8F", "ref": "495057",
    "plan_head": "4A90E2", "plan_fill": "EAF3FF",
    "fact_head": "2E9E5B", "fact_fill": "E6F6EC",
    "input_fill": "FFFBEA", "band": "F7F7FB", "grid": "D9DCE3",
    "text": "2B2D42", "muted": "8D99AE", "pos": "1B9E5A", "neg": "D62839",
    "done_fill": "EFEFEF", "today": "FFF3BF",
}

FONT = "Arial"
MONEY = '#,##0 "₸";[Red]-#,##0 "₸";"–"'
PCT = '0%;-0%;"–"'

thin = Side(style="thin", color=C["grid"])
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def font(size=10, bold=False, color=None, italic=False):
    return Font(name=FONT, size=size, bold=bold, color=color or C["text"], italic=italic)


def fill(hex_color):
    return PatternFill("solid", start_color=hex_color, end_color=hex_color)


def q(sheet):
    return f"'{sheet}'"


def title(ws, text, subtitle, color, width_cols):
    last = get_column_letter(width_cols)
    ws.merge_cells(f"A1:{last}1")
    ws.merge_cells(f"A2:{last}2")
    ws["A1"] = text
    ws["A1"].font = Font(name=FONT, size=18, bold=True, color="FFFFFF")
    ws["A1"].fill = fill(color)
    ws["A1"].alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[1].height = 36
    ws["A2"] = subtitle
    ws["A2"].font = font(9, italic=True, color=C["muted"])
    ws["A2"].alignment = Alignment(vertical="center", indent=1, wrap_text=True)
    ws.row_dimensions[2].height = 30
    for col in range(1, width_cols + 1):
        ws.cell(1, col).fill = fill(color)
    ws.sheet_properties.tabColor = color
    ws.sheet_view.showGridLines = False


def header(ws, row, headers, color, kinds=None):
    """kinds: список 'plan' / 'fact' / None для подсветки колонок план/факт."""
    for i, h in enumerate(headers, start=1):
        cell = ws.cell(row, i, h)
        kind = kinds[i - 1] if kinds else None
        bg = C["plan_head"] if kind == "plan" else C["fact_head"] if kind == "fact" else color
        cell.fill = fill(bg)
        cell.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[row].height = 34


def widths(ws, values):
    for i, w in enumerate(values, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def style_range(ws, r1, r2, c1, c2, *, fmt=None, bg=None, align=None, fnt=None):
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            cell = ws.cell(r, c)
            cell.border = BORDER
            cell.font = fnt or font()
            if fmt:
                cell.number_format = fmt
            if bg:
                cell.fill = fill(bg)
            if align:
                cell.alignment = Alignment(horizontal=align, vertical="center")


def add_name(wb, name, sheet, ref):
    wb.defined_names[name] = DefinedName(name, attr_text=f"{q(sheet)}!{ref}")


def add_list_validation(ws, rng, source, prompt=None):
    dv = DataValidation(type="list", formula1=f"={source}", allow_blank=True)
    dv.error = "Выберите значение из списка (его можно дополнить в листе «📚 Справочник»)"
    dv.errorTitle = "Значение не из справочника"
    dv.showErrorMessage = True
    if prompt:
        dv.prompt = prompt
        dv.showInputMessage = True
    ws.add_data_validation(dv)
    dv.add(rng)


def line_chart(title_text, y_title, height=8, width=22):
    ch = LineChart()
    ch.title = title_text
    ch.style = 2
    ch.height = height
    ch.width = width
    ch.y_axis.title = y_title
    ch.y_axis.number_format = '#,##0'
    ch.y_axis.majorGridlines.spPr = GraphicalProperties(ln=None)
    ch.legend.position = "b"
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    return ch


def color_series(ch, colors, dashed=None):
    dashed = dashed or []
    for i, s in enumerate(ch.series):
        s.graphicalProperties.line.solidFill = colors[i]
        s.graphicalProperties.line.width = 28000
        s.smooth = False
        if i in dashed:
            s.graphicalProperties.line.dashStyle = "dash"


wb = Workbook()
home = wb.active
home.title = S_HOME
exp = wb.create_sheet(S_EXP)
inc = wb.create_sheet(S_INC)
mon = wb.create_sheet(S_MON)
yr = wb.create_sheet(S_YEAR)
wish = wb.create_sheet(S_WISH)
ref = wb.create_sheet(S_REF)

# ---------------------------------------------------------------- СПРАВОЧНИК
title(ref, "📚 Справочник",
      "Здесь живут все выпадающие списки. Добавляйте свои статьи в пустые жёлтые строки — "
      "они сразу появятся в списках листов «Расходы» и «Доходы».", C["ref"], 17)
widths(ref, [10, 14, 3, 24, 13, 34, 3, 30, 3, 22, 3, 30, 3, 18, 18, 3, 16])

CATS = [
    ("🛒 Нужда", "Нет", "Еда, жильё, транспорт — без этого никак"),
    ("🧾 Обязательный платёж", "Нет", "Кредиты, налоги, страховки, подписки"),
    ("🍭 Хотелка", "Да", "Маленькие радости «хочу прямо сейчас»"),
    ("💫 Желание", "Да", "Крупнее хотелки, стоит накопить"),
    ("🌟 Мечта", "Да", "Большая цель: путешествие, машина, квартира"),
    ("🎁 Подарок", "Нет", "Подарки близким"),
    ("🏦 Накопление", "Нет", "Подушка безопасности, депозит"),
    ("📈 Инвестиция", "Нет", "Акции, бизнес, активы"),
    ("❤️ Здоровье", "Нет", "Врачи, спорт, лекарства"),
    ("🎓 Образование", "Нет", "Курсы, книги, обучение"),
    ("🎉 Развлечения", "Нет", "Кафе, кино, отдых"),
]
EXP_ITEMS = ["Продукты", "Аренда / ипотека", "Коммунальные услуги", "Транспорт / такси",
             "Связь и интернет", "Одежда", "Кафе и рестораны", "Путешествия", "Техника",
             "Подушка безопасности", "Подарки", "Спорт", "Курсы"]
INC_TYPES = ["💼 Зарплата", "🚀 Проект", "🛠 Разовая услуга", "📈 Инвестиции",
             "🏆 Премия / бонус", "🏠 Аренда", "🎁 Подарок", "💡 Прочее"]
INC_ITEMS = ["Основная работа", "Фриланс-проект", "Консультация", "Дивиденды", "Кэшбэк",
             "Подработка"]
EXP_STATUS = ["⏳ Запланировано", "✅ Выполнено", "⏸ Отложено", "❌ Отменено"]
INC_STATUS = ["⏳ Ожидается", "✅ Получено", "❌ Не получено"]
PRIORITY = ["🔥 Высокий", "⭐ Средний", "💤 Низкий"]

R0 = 4  # строка заголовков справочника


def ref_list(col, head, values, rows, name, editable=False):
    header_cell = ref.cell(R0, col, head)
    header_cell.fill = fill(C["ref"])
    header_cell.font = Font(name=FONT, bold=True, color="FFFFFF")
    header_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    header_cell.border = BORDER
    for i in range(rows):
        cell = ref.cell(R0 + 1 + i, col, values[i] if i < len(values) else None)
        cell.border = BORDER
        cell.font = font(color="0000FF" if editable else C["text"])
        if editable:
            cell.fill = fill(C["input_fill"])
    col_l = get_column_letter(col)
    add_name(wb, name, S_REF, f"${col_l}${R0 + 1}:${col_l}${R0 + rows}")


ref.row_dimensions[R0].height = 34
ref_list(1, "📆 Год", YEARS, 5, "lst_years")
ref_list(2, "🗓 Месяц", MONTHS, 12, "lst_months")

# Категории расходов: 3 колонки
CAT_ROWS = 25
for col, head in ((4, "🏷 Категория расхода"), (5, "✨ В список хотелок?"), (6, "💬 Пояснение")):
    c = ref.cell(R0, col, head)
    c.fill = fill(C["ref"])
    c.font = Font(name=FONT, bold=True, color="FFFFFF")
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = BORDER
for i in range(CAT_ROWS):
    r = R0 + 1 + i
    vals = CATS[i] if i < len(CATS) else (None, None, None)
    for j, v in enumerate(vals):
        cell = ref.cell(r, 4 + j, v)
        cell.border = BORDER
        cell.fill = fill(C["input_fill"])
        cell.font = font(color="0000FF" if j < 2 else C["muted"], italic=(j == 2))
add_name(wb, "lst_cats", S_REF, f"$D${R0 + 1}:$D${R0 + CAT_ROWS}")
add_name(wb, "lst_cat_wish", S_REF, f"$E${R0 + 1}:$E${R0 + CAT_ROWS}")
add_list_validation(ref, f"E{R0 + 1}:E{R0 + CAT_ROWS}", '"Да,Нет"')
ref.cell(R0 - 1, 4, "Категория с «Да» автоматически попадает в лист «✨ Хотелки»").font = \
    font(8, italic=True, color=C["muted"])

ref_list(8, "📝 Статьи расходов", EXP_ITEMS, 100, "lst_exp_items", editable=True)
ref_list(10, "💼 Тип дохода", INC_TYPES, 20, "lst_inc_types", editable=True)
ref_list(12, "💰 Источники дохода", INC_ITEMS, 100, "lst_inc_items", editable=True)
ref_list(14, "🚦 Статус расхода", EXP_STATUS, 4, "lst_exp_status")
ref_list(15, "🚦 Статус дохода", INC_STATUS, 3, "lst_inc_status")
ref_list(17, "🎯 Приоритет", PRIORITY, 3, "lst_priority")
ref.freeze_panes = "A5"

DONE_EXP = EXP_STATUS[1]
CANCEL_EXP = EXP_STATUS[3]
DONE_INC = INC_STATUS[1]
FAIL_INC = INC_STATUS[2]

# ---------------------------------------------------------------- РАСХОДЫ
EXP_HEAD = ["📆 Год", "🗓 Месяц", "📝 Статья", "🏷 Категория", "✏️ Описание", "🎯 Приоритет",
            "📝 ПЛАН, ₸", "✅ ФАКТ, ₸", "± Разница\n(план − факт)", "🚦 Статус", "💬 Комментарий",
            "служ. №хотелки"]
EXP_KINDS = [None] * 6 + ["plan", "fact", None, None, None, None]
title(exp, "💸 Расходы — план и факт",
      "Одна строка = одна трата. Синяя колонка «ПЛАН» — сколько собираетесь потратить, зелёная «ФАКТ» — "
      "сколько реально ушло. Статус «✅ Выполнено» зачёркивает строку. Год, месяц, статья, категория — из списков.",
      C["exp"], 11)
header(exp, 4, EXP_HEAD, C["exp"], EXP_KINDS)
widths(exp, [9, 12, 22, 22, 32, 13, 14, 14, 14, 17, 26, 6])
exp.column_dimensions["L"].hidden = True

EXP_EXAMPLES = [
    (2026, "Октябрь", "Продукты", "🛒 Нужда", "Продукты на месяц", "🔥 Высокий", 120000, 131500, DONE_EXP),
    (2026, "Октябрь", "Аренда / ипотека", "🧾 Обязательный платёж", "Аренда квартиры", "🔥 Высокий", 250000, 250000, DONE_EXP),
    (2026, "Октябрь", "Подушка безопасности", "🏦 Накопление", "10% от дохода", "⭐ Средний", 60000, 60000, DONE_EXP),
    (2026, "Октябрь", "Техника", "🍭 Хотелка", "Беспроводные наушники", "⭐ Средний", 45000, 42000, DONE_EXP),
    (2026, "Ноябрь", "Продукты", "🛒 Нужда", "Продукты на месяц", "🔥 Высокий", 120000, None, EXP_STATUS[0]),
    (2026, "Ноябрь", "Аренда / ипотека", "🧾 Обязательный платёж", "Аренда квартиры", "🔥 Высокий", 250000, None, EXP_STATUS[0]),
    (2026, "Декабрь", "Подарки", "🎁 Подарок", "Новогодние подарки семье", "⭐ Средний", 80000, None, EXP_STATUS[0]),
    (2027, "Март", "Одежда", "💫 Желание", "Новое пальто", "💤 Низкий", 90000, None, EXP_STATUS[0]),
    (2027, "Июль", "Путешествия", "🌟 Мечта", "Отпуск в Грузии", "🔥 Высокий", 600000, None, EXP_STATUS[0]),
    (2028, "Май", "Курсы", "🎓 Образование", "Курс английского", "⭐ Средний", 150000, None, EXP_STATUS[0]),
]

for r in range(EXP_FIRST, EXP_LAST + 1):
    ex = EXP_EXAMPLES[r - EXP_FIRST] if r - EXP_FIRST < len(EXP_EXAMPLES) else None
    if ex:
        for j, v in enumerate(ex[:8]):
            exp.cell(r, j + 1, v)
        exp.cell(r, 10, ex[8])
    exp.cell(r, 9, f'=IF(OR(G{r}="",H{r}=""),"",G{r}-H{r})')
    exp.cell(r, 12, f'=IF(IFERROR(INDEX(lst_cat_wish,MATCH(D{r},lst_cats,0))="Да",FALSE),MAX($L$4:L{r - 1})+1,"")')
    for c in range(1, 13):
        cell = exp.cell(r, c)
        cell.border = BORDER
        cell.font = font(color="0000FF" if c not in (9, 12) else C["text"])
        cell.alignment = Alignment(vertical="center", horizontal="center" if c in (1, 2, 6) else None)
    exp.cell(r, 7).fill = fill(C["plan_fill"])
    exp.cell(r, 8).fill = fill(C["fact_fill"])
    for c in (7, 8, 9):
        exp.cell(r, c).number_format = MONEY

rng = lambda col: f"{col}{EXP_FIRST}:{col}{EXP_LAST}"  # noqa: E731
add_list_validation(exp, rng("A"), "lst_years")
add_list_validation(exp, rng("B"), "lst_months")
add_list_validation(exp, rng("C"), "lst_exp_items")
add_list_validation(exp, rng("D"), "lst_cats")
add_list_validation(exp, rng("F"), "lst_priority")
add_list_validation(exp, rng("J"), "lst_exp_status")

for nm, col in (("exp_year", "A"), ("exp_month", "B"), ("exp_item", "C"), ("exp_cat", "D"),
                ("exp_desc", "E"), ("exp_prio", "F"), ("exp_plan", "G"), ("exp_fact", "H"),
                ("exp_status", "J"), ("exp_wishno", "L")):
    add_name(wb, nm, S_EXP, f"${col}${EXP_FIRST}:${col}${EXP_LAST}")

data = f"A{EXP_FIRST}:K{EXP_LAST}"
exp.conditional_formatting.add(data, FormulaRule(
    formula=[f'$J{EXP_FIRST}="{DONE_EXP}"'], stopIfTrue=True,
    font=Font(strike=True, color="8D99AE"), fill=fill(C["done_fill"])))
exp.conditional_formatting.add(data, FormulaRule(
    formula=[f'$J{EXP_FIRST}="{CANCEL_EXP}"'], stopIfTrue=True,
    font=Font(strike=True, italic=True, color="C0C0C0")))
exp.conditional_formatting.add(f"H{EXP_FIRST}:H{EXP_LAST}", FormulaRule(
    formula=[f'AND($H{EXP_FIRST}<>"",$G{EXP_FIRST}<>"",$H{EXP_FIRST}>$G{EXP_FIRST})'],
    font=Font(bold=True, color=C["neg"])))
exp.conditional_formatting.add(f"I{EXP_FIRST}:I{EXP_LAST}", CellIsRule(
    operator="lessThan", formula=["0"], font=Font(bold=True, color=C["neg"])))
exp.conditional_formatting.add(f"I{EXP_FIRST}:I{EXP_LAST}", CellIsRule(
    operator="greaterThan", formula=["0"], font=Font(bold=True, color=C["pos"])))
exp.conditional_formatting.add(f"F{EXP_FIRST}:F{EXP_LAST}", FormulaRule(
    formula=[f'$F{EXP_FIRST}="{PRIORITY[0]}"'], fill=fill("FFE3E3")))
exp.freeze_panes = "C5"
exp.auto_filter.ref = f"A4:K{EXP_LAST}"
exp["G3"] = f"=SUBTOTAL(9,G{EXP_FIRST}:G{EXP_LAST})"
exp["H3"] = f"=SUBTOTAL(9,H{EXP_FIRST}:H{EXP_LAST})"
exp["F3"] = "Σ видимых →"
exp["F3"].font = font(9, True, C["muted"])
exp["F3"].alignment = Alignment(horizontal="right")
for c in ("G3", "H3"):
    exp[c].number_format = MONEY
    exp[c].font = font(10, True)
exp["G3"].fill = fill(C["plan_fill"])
exp["H3"].fill = fill(C["fact_fill"])
exp["G4"].comment = Comment("Сумма, которую ПЛАНИРУЕТЕ потратить. Заполняется заранее.", "Планер")
exp["H4"].comment = Comment("Сколько РЕАЛЬНО потратили. Заполняется по факту.", "Планер")

# ---------------------------------------------------------------- ДОХОДЫ
INC_HEAD = ["📆 Год", "🗓 Месяц", "💰 Источник", "💼 Тип дохода", "✏️ Описание",
            "📝 ПЛАН, ₸\n(потенциально)", "✅ ФАКТ, ₸\n(получено)", "± Разница\n(факт − план)",
            "🚦 Статус", "💬 Комментарий"]
INC_KINDS = [None] * 5 + ["plan", "fact", None, None, None]
title(inc, "💰 Доходы — сколько могу заработать и сколько получил",
      "Синяя колонка «ПЛАН» — сколько потенциально можно заработать (ЗП, проект, разовая услуга…), "
      "зелёная «ФАКТ» — сколько реально пришло. Статус «✅ Получено» зачёркивает строку.",
      C["inc"], 10)
header(inc, 4, INC_HEAD, C["inc"], INC_KINDS)
widths(inc, [9, 12, 22, 20, 32, 16, 16, 14, 16, 26])

INC_EXAMPLES = [
    (2026, "Октябрь", "Основная работа", "💼 Зарплата", "Оклад", 500000, 500000, DONE_INC),
    (2026, "Октябрь", "Консультация", "🛠 Разовая услуга", "Настройка CRM", 50000, 65000, DONE_INC),
    (2026, "Ноябрь", "Основная работа", "💼 Зарплата", "Оклад", 500000, None, INC_STATUS[0]),
    (2026, "Ноябрь", "Фриланс-проект", "🚀 Проект", "Сайт для клиента (50% аванс)", 300000, None, INC_STATUS[0]),
    (2026, "Декабрь", "Основная работа", "💼 Зарплата", "Оклад + годовая премия", 750000, None, INC_STATUS[0]),
    (2027, "Январь", "Основная работа", "💼 Зарплата", "Оклад (после индексации)", 550000, None, INC_STATUS[0]),
]
for r in range(INC_FIRST, INC_LAST + 1):
    ex = INC_EXAMPLES[r - INC_FIRST] if r - INC_FIRST < len(INC_EXAMPLES) else None
    if ex:
        for j, v in enumerate(ex[:7]):
            inc.cell(r, j + 1, v)
        inc.cell(r, 9, ex[7])
    inc.cell(r, 8, f'=IF(OR(F{r}="",G{r}=""),"",G{r}-F{r})')
    for c in range(1, 11):
        cell = inc.cell(r, c)
        cell.border = BORDER
        cell.font = font(color="0000FF" if c != 8 else C["text"])
        cell.alignment = Alignment(vertical="center", horizontal="center" if c in (1, 2) else None)
    inc.cell(r, 6).fill = fill(C["plan_fill"])
    inc.cell(r, 7).fill = fill(C["fact_fill"])
    for c in (6, 7, 8):
        inc.cell(r, c).number_format = MONEY

rng = lambda col: f"{col}{INC_FIRST}:{col}{INC_LAST}"  # noqa: E731
add_list_validation(inc, rng("A"), "lst_years")
add_list_validation(inc, rng("B"), "lst_months")
add_list_validation(inc, rng("C"), "lst_inc_items")
add_list_validation(inc, rng("D"), "lst_inc_types")
add_list_validation(inc, rng("I"), "lst_inc_status")
for nm, col in (("inc_year", "A"), ("inc_month", "B"), ("inc_type", "D"), ("inc_plan", "F"),
                ("inc_fact", "G"), ("inc_status", "I")):
    add_name(wb, nm, S_INC, f"${col}${INC_FIRST}:${col}${INC_LAST}")

data = f"A{INC_FIRST}:J{INC_LAST}"
inc.conditional_formatting.add(data, FormulaRule(
    formula=[f'$I{INC_FIRST}="{DONE_INC}"'], stopIfTrue=True,
    font=Font(strike=True, color="8D99AE"), fill=fill(C["done_fill"])))
inc.conditional_formatting.add(data, FormulaRule(
    formula=[f'$I{INC_FIRST}="{FAIL_INC}"'], stopIfTrue=True,
    font=Font(strike=True, italic=True, color=C["neg"])))
inc.conditional_formatting.add(f"H{INC_FIRST}:H{INC_LAST}", CellIsRule(
    operator="lessThan", formula=["0"], font=Font(bold=True, color=C["neg"])))
inc.conditional_formatting.add(f"H{INC_FIRST}:H{INC_LAST}", CellIsRule(
    operator="greaterThan", formula=["0"], font=Font(bold=True, color=C["pos"])))
inc.freeze_panes = "C5"
inc.auto_filter.ref = f"A4:J{INC_LAST}"
inc["E3"] = "Σ видимых →"
inc["E3"].font = font(9, True, C["muted"])
inc["E3"].alignment = Alignment(horizontal="right")
inc["F3"] = f"=SUBTOTAL(9,F{INC_FIRST}:F{INC_LAST})"
inc["G3"] = f"=SUBTOTAL(9,G{INC_FIRST}:G{INC_LAST})"
for c in ("F3", "G3"):
    inc[c].number_format = MONEY
    inc[c].font = font(10, True)
inc["F3"].fill = fill(C["plan_fill"])
inc["G3"].fill = fill(C["fact_fill"])

# ---------------------------------------------------------------- ПО МЕСЯЦАМ
MON_HEAD = ["📆 Год", "🗓 Месяц", "Подпись", "💰 Доход\nПЛАН", "💰 Доход\nФАКТ", "% дохода\nполучено",
            "💸 Расход\nПЛАН", "💸 Расход\nФАКТ", "% бюджета\nпотрачено", "⚖️ Остаток\nПЛАН",
            "⚖️ Остаток\nФАКТ", "🏦 Накоплено\nПЛАН (нараст.)", "🏦 Накоплено\nФАКТ (нараст.)",
            "🏦 Отложено в\n«Накопление» (факт)"]
MON_KINDS = [None, None, None, "plan", "fact", None, "plan", "fact", None, "plan", "fact", "plan", "fact", "fact"]
title(mon, "📅 Сводка по месяцам: январь 2026 → декабрь 2030",
      "Считается автоматически из листов «Расходы» и «Доходы» — здесь ничего вводить не нужно. "
      "Текущий месяц подсвечен жёлтым. Отменённые траты в план не входят.", C["mon"], 14)
header(mon, 4, MON_HEAD, C["mon"], MON_KINDS)
widths(mon, [8, 12, 9, 14, 14, 11, 14, 14, 11, 14, 14, 16, 16, 16])
mon.column_dimensions["C"].hidden = True
M_FIRST = 5
r = M_FIRST
for y in YEARS:
    for m in MONTHS:
        mon.cell(r, 1, y)
        mon.cell(r, 2, m)
        mon.cell(r, 3, f'=LEFT(B{r},3)&" "&RIGHT(A{r},2)')
        mon.cell(r, 4, f'=SUMIFS(inc_plan,inc_year,A{r},inc_month,B{r},inc_status,"<>{FAIL_INC}")')
        mon.cell(r, 5, f"=SUMIFS(inc_fact,inc_year,A{r},inc_month,B{r})")
        mon.cell(r, 6, f'=IF(D{r}=0,"",E{r}/D{r})')
        mon.cell(r, 7, f'=SUMIFS(exp_plan,exp_year,A{r},exp_month,B{r},exp_status,"<>{CANCEL_EXP}")')
        mon.cell(r, 8, f"=SUMIFS(exp_fact,exp_year,A{r},exp_month,B{r})")
        mon.cell(r, 9, f'=IF(G{r}=0,"",H{r}/G{r})')
        mon.cell(r, 10, f"=D{r}-G{r}")
        mon.cell(r, 11, f"=E{r}-H{r}")
        mon.cell(r, 12, f"=J{r}" if r == M_FIRST else f"=L{r - 1}+J{r}")
        mon.cell(r, 13, f"=K{r}" if r == M_FIRST else f"=M{r - 1}+K{r}")
        mon.cell(r, 14, f'=SUMIFS(exp_fact,exp_year,A{r},exp_month,B{r},exp_cat,"{CATS[6][0]}")')
        r += 1
M_LAST = r - 1
style_range(mon, M_FIRST, M_LAST, 1, 14)
for rr in range(M_FIRST, M_LAST + 1):
    for c in (4, 5, 7, 8, 10, 11, 12, 13, 14):
        mon.cell(rr, c).number_format = MONEY
    for c in (6, 9):
        mon.cell(rr, c).number_format = PCT
    mon.cell(rr, 1).alignment = Alignment(horizontal="center")
    for c, kind in enumerate(MON_KINDS, start=1):
        if kind == "plan":
            mon.cell(rr, c).fill = fill(C["plan_fill"])
        elif kind == "fact":
            mon.cell(rr, c).fill = fill(C["fact_fill"])
    if rr != M_FIRST and mon.cell(rr, 2).value == "Январь":  # разделитель лет
        for c in range(1, 15):
            mon.cell(rr, c).border = Border(left=thin, right=thin, bottom=thin,
                                            top=Side(style="medium", color=C["mon"]))
# итог
mon.cell(M_LAST + 1, 1, "Σ ИТОГО")
mon.merge_cells(start_row=M_LAST + 1, start_column=1, end_row=M_LAST + 1, end_column=3)
for c in (4, 5, 7, 8, 10, 11, 14):
    L = get_column_letter(c)
    mon.cell(M_LAST + 1, c, f"=SUM({L}{M_FIRST}:{L}{M_LAST})")
mon.cell(M_LAST + 1, 6, f'=IF(D{M_LAST + 1}=0,"",E{M_LAST + 1}/D{M_LAST + 1})')
mon.cell(M_LAST + 1, 9, f'=IF(G{M_LAST + 1}=0,"",H{M_LAST + 1}/G{M_LAST + 1})')
style_range(mon, M_LAST + 1, M_LAST + 1, 1, 14, bg="E3ECFF", fnt=font(10, True))
for c in (4, 5, 7, 8, 10, 11, 14):
    mon.cell(M_LAST + 1, c).number_format = MONEY
for c in (6, 9):
    mon.cell(M_LAST + 1, c).number_format = PCT

mrange = f"A{M_FIRST}:N{M_LAST}"
mon.conditional_formatting.add(mrange, FormulaRule(
    formula=[f"AND($A{M_FIRST}=YEAR(TODAY()),MATCH($B{M_FIRST},lst_months,0)=MONTH(TODAY()))"],
    fill=fill(C["today"]), font=Font(bold=True)))
for col in ("J", "K", "L", "M"):
    mon.conditional_formatting.add(f"{col}{M_FIRST}:{col}{M_LAST}", CellIsRule(
        operator="lessThan", formula=["0"], font=Font(bold=True, color=C["neg"])))
    mon.conditional_formatting.add(f"{col}{M_FIRST}:{col}{M_LAST}", CellIsRule(
        operator="greaterThan", formula=["0"], font=Font(color=C["pos"])))
mon.conditional_formatting.add(f"F{M_FIRST}:F{M_LAST}", DataBarRule(
    start_type="num", start_value=0, end_type="num", end_value=1, color="2E9E5B"))
mon.conditional_formatting.add(f"I{M_FIRST}:I{M_LAST}", CellIsRule(
    operator="greaterThan", formula=["1"], font=Font(bold=True, color=C["neg"]), fill=fill("FFE3E3")))
mon.conditional_formatting.add(f"I{M_FIRST}:I{M_LAST}", DataBarRule(
    start_type="num", start_value=0, end_type="num", end_value=1, color="E76F51"))
mon.freeze_panes = "D5"

cats_ref = Reference(mon, min_col=3, min_row=M_FIRST, max_row=M_LAST)
ch = line_chart("💰 Доход по месяцам: план vs факт", "₸", height=9, width=26)
ch.add_data(Reference(mon, min_col=4, max_col=5, min_row=4, max_row=M_LAST), titles_from_data=True)
ch.set_categories(cats_ref)
color_series(ch, [C["plan_head"], C["fact_head"]], dashed=[0])
mon.add_chart(ch, "P4")

ch = line_chart("💸 Расход по месяцам: план vs факт", "₸", height=9, width=26)
ch.add_data(Reference(mon, min_col=7, max_col=8, min_row=4, max_row=M_LAST), titles_from_data=True)
ch.set_categories(cats_ref)
color_series(ch, ["F4A261", C["neg"]], dashed=[0])
mon.add_chart(ch, "P23")

ch = line_chart("🏦 Накопленный остаток: план vs факт", "₸", height=9, width=26)
ch.add_data(Reference(mon, min_col=12, max_col=13, min_row=4, max_row=M_LAST), titles_from_data=True)
ch.set_categories(cats_ref)
color_series(ch, [C["year"], C["inc"]], dashed=[0])
mon.add_chart(ch, "P42")

# ---------------------------------------------------------------- ПО ГОДАМ
YEAR_HEAD = ["📆 Год", "💰 Доход\nПЛАН", "💰 Доход\nФАКТ", "% дохода\nполучено", "💸 Расход\nПЛАН",
             "💸 Расход\nФАКТ", "% бюджета\nпотрачено", "⚖️ Остаток\nПЛАН", "⚖️ Остаток\nФАКТ",
             "💎 Норма сбере-\nжений ПЛАН", "💎 Норма сбере-\nжений ФАКТ"]
YEAR_KINDS = [None, "plan", "fact", None, "plan", "fact", None, "plan", "fact", "plan", "fact"]
title(yr, "📊 Сводка по годам 2026–2030",
      "Верхняя таблица — общий итог по годам, нижняя — расходы в разрезе категорий. "
      "Всё считается автоматически. Норма сбережений = остаток ÷ доход.", C["year"], 13)
header(yr, 4, YEAR_HEAD, C["year"], YEAR_KINDS)
widths(yr, [12, 15, 15, 11, 15, 15, 11, 15, 15, 13, 13, 15, 15])
Y_FIRST = 5
for i, y in enumerate(YEARS):
    r = Y_FIRST + i
    yr.cell(r, 1, y)
    yr.cell(r, 2, f'=SUMIFS(inc_plan,inc_year,A{r},inc_status,"<>{FAIL_INC}")')
    yr.cell(r, 3, f"=SUMIFS(inc_fact,inc_year,A{r})")
    yr.cell(r, 4, f'=IF(B{r}=0,"",C{r}/B{r})')
    yr.cell(r, 5, f'=SUMIFS(exp_plan,exp_year,A{r},exp_status,"<>{CANCEL_EXP}")')
    yr.cell(r, 6, f"=SUMIFS(exp_fact,exp_year,A{r})")
    yr.cell(r, 7, f'=IF(E{r}=0,"",F{r}/E{r})')
    yr.cell(r, 8, f"=B{r}-E{r}")
    yr.cell(r, 9, f"=C{r}-F{r}")
    yr.cell(r, 10, f'=IF(B{r}=0,"",H{r}/B{r})')
    yr.cell(r, 11, f'=IF(C{r}=0,"",I{r}/C{r})')
Y_LAST = Y_FIRST + len(YEARS) - 1
YT = Y_LAST + 1
yr.cell(YT, 1, "Σ 2026–2030")
for c in (2, 3, 5, 6, 8, 9):
    L = get_column_letter(c)
    yr.cell(YT, c, f"=SUM({L}{Y_FIRST}:{L}{Y_LAST})")
yr.cell(YT, 4, f'=IF(B{YT}=0,"",C{YT}/B{YT})')
yr.cell(YT, 7, f'=IF(E{YT}=0,"",F{YT}/E{YT})')
yr.cell(YT, 10, f'=IF(B{YT}=0,"",H{YT}/B{YT})')
yr.cell(YT, 11, f'=IF(C{YT}=0,"",I{YT}/C{YT})')
style_range(yr, Y_FIRST, YT, 1, 11)
for r in range(Y_FIRST, YT + 1):
    yr.cell(r, 1).alignment = Alignment(horizontal="center")
    yr.cell(r, 1).font = font(11, True)
    for c in (2, 3, 5, 6, 8, 9):
        yr.cell(r, c).number_format = MONEY
    for c in (4, 7, 10, 11):
        yr.cell(r, c).number_format = PCT
    for c, kind in enumerate(YEAR_KINDS, start=1):
        if kind == "plan":
            yr.cell(r, c).fill = fill(C["plan_fill"])
        elif kind == "fact":
            yr.cell(r, c).fill = fill(C["fact_fill"])
    yr.row_dimensions[r].height = 22
for c in range(1, 12):
    yr.cell(YT, c).fill = fill("EDE3FF")
    yr.cell(YT, c).font = font(10, True)
yr.conditional_formatting.add(f"A{Y_FIRST}:K{Y_LAST}", FormulaRule(
    formula=[f"$A{Y_FIRST}=YEAR(TODAY())"], font=Font(bold=True, color="8338EC")))
for col in ("H", "I"):
    yr.conditional_formatting.add(f"{col}{Y_FIRST}:{col}{YT}", CellIsRule(
        operator="lessThan", formula=["0"], font=Font(bold=True, color=C["neg"])))
for col in ("D", "J", "K"):
    yr.conditional_formatting.add(f"{col}{Y_FIRST}:{col}{YT}", DataBarRule(
        start_type="num", start_value=0, end_type="num", end_value=1, color="8338EC"))
yr.conditional_formatting.add(f"G{Y_FIRST}:G{YT}", CellIsRule(
    operator="greaterThan", formula=["1"], font=Font(bold=True, color=C["neg"]), fill=fill("FFE3E3")))

# Категории × годы
CR = YT + 3
yr.cell(CR - 1, 1, "🏷 Расходы по категориям и годам").font = font(13, True, C["year"])
yr.merge_cells(start_row=CR, start_column=1, end_row=CR + 1, end_column=1)
yr.cell(CR, 1, "Категория")
col = 2
for y in YEARS + ["Σ Итого"]:
    yr.merge_cells(start_row=CR, start_column=col, end_row=CR, end_column=col + 1)
    yr.cell(CR, col, str(y))
    yr.cell(CR + 1, col, "📝 План")
    yr.cell(CR + 1, col + 1, "✅ Факт")
    col += 2
for rr in (CR, CR + 1):
    for c in range(1, 14):
        cell = yr.cell(rr, c)
        sub = rr == CR + 1 and c > 1
        bg = (C["plan_head"] if c % 2 == 0 else C["fact_head"]) if sub else C["year"]
        cell.fill = fill(bg)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER
C_FIRST = CR + 2
N_CAT = 15
for i in range(N_CAT):
    r = C_FIRST + i
    src = R0 + 1 + i
    yr.cell(r, 1, f"=IF({q(S_REF)}!D{src}=\"\",\"\",{q(S_REF)}!D{src})")
    col = 2
    for y in YEARS:
        yr.cell(r, col, f'=IF($A{r}="","",SUMIFS(exp_plan,exp_cat,$A{r},exp_year,{y},exp_status,"<>{CANCEL_EXP}"))')
        yr.cell(r, col + 1, f'=IF($A{r}="","",SUMIFS(exp_fact,exp_cat,$A{r},exp_year,{y}))')
        col += 2
    plan_cells = ",".join(f"{get_column_letter(2 + 2 * k)}{r}" for k in range(len(YEARS)))
    fact_cells = ",".join(f"{get_column_letter(3 + 2 * k)}{r}" for k in range(len(YEARS)))
    yr.cell(r, 12, f'=IF($A{r}="","",SUM({plan_cells}))')
    yr.cell(r, 13, f'=IF($A{r}="","",SUM({fact_cells}))')
C_LAST = C_FIRST + N_CAT - 1
style_range(yr, C_FIRST, C_LAST, 1, 13, fmt=MONEY)
for r in range(C_FIRST, C_LAST + 1):
    yr.cell(r, 1).number_format = "General"
    for c in range(2, 14):
        yr.cell(r, c).fill = fill(C["plan_fill"] if c % 2 == 0 else C["fact_fill"])
    for c in (12, 13):
        yr.cell(r, c).font = font(10, True)
CT = C_LAST + 1
yr.cell(CT, 1, "Σ ИТОГО")
for c in range(2, 14):
    L = get_column_letter(c)
    yr.cell(CT, c, f"=SUM({L}{C_FIRST}:{L}{C_LAST})")
style_range(yr, CT, CT, 1, 13, fmt=MONEY, bg="EDE3FF", fnt=font(10, True))
yr.conditional_formatting.add(f"M{C_FIRST}:M{C_LAST}", DataBarRule(
    start_type="min", end_type="max", color="FF5D8F"))
yr.freeze_panes = "B5"

ych_cats = Reference(yr, min_col=1, min_row=Y_FIRST, max_row=Y_LAST)
ch = line_chart("📈 Доход по годам: план vs факт", "₸", height=8, width=18)
ch.add_data(Reference(yr, min_col=2, max_col=3, min_row=4, max_row=Y_LAST), titles_from_data=True)
ch.set_categories(ych_cats)
color_series(ch, [C["plan_head"], C["fact_head"]], dashed=[0])
yr.add_chart(ch, "O4")
ch = line_chart("📉 Расход по годам: план vs факт", "₸", height=8, width=18)
ch.add_data(Reference(yr, min_col=5, max_col=6, min_row=4, max_row=Y_LAST), titles_from_data=True)
ch.set_categories(ych_cats)
color_series(ch, ["F4A261", C["neg"]], dashed=[0])
yr.add_chart(ch, "O21")

# ---------------------------------------------------------------- ХОТЕЛКИ
WISH_HEAD = ["№", "✨ Хотелка", "🏷 Категория", "🎯 Приоритет", "📆 Год", "🗓 Месяц",
             "📝 ПЛАН, ₸", "✅ ФАКТ, ₸", "🚦 Статус"]
title(wish, "✨ Список хотелок, желаний и мечт",
      "Собирается АВТОМАТИЧЕСКИ из листа «💸 Расходы» по категориям с отметкой «Да» в справочнике. "
      "Чтобы вычеркнуть хотелку — поставьте ей статус «✅ Выполнено» в листе «Расходы».", C["wish"], 9)
header(wish, 6, WISH_HEAD, C["wish"], [None] * 6 + ["plan", "fact", None])
widths(wish, [6, 36, 18, 14, 9, 12, 14, 14, 18])
W_FIRST = 7
W_LAST = W_FIRST + WISH_ROWS - 1
# мини-статистика
stats = [("Всего хотелок", "=COUNT(exp_wishno)", "0"),
         ("✅ Исполнено", f'=COUNTIFS(exp_wishno,">0",exp_status,"{DONE_EXP}")', "0"),
         ("🎯 Прогресс", '=IF(B4=0,"",C4/B4)', PCT),
         ("💰 Ещё нужно, ₸", f'=SUMPRODUCT((exp_wishno<>"")*(exp_status<>"{DONE_EXP}")*(exp_status<>"{CANCEL_EXP}"),exp_plan)', MONEY)]
# раскладываем статистику: подписи в строке 3, значения в строке 4
for i, (lbl, f, fmt) in enumerate(stats):
    col = 2 + i
    c1 = wish.cell(3, col, lbl)
    c1.font = font(9, True, C["muted"])
    c1.alignment = Alignment(horizontal="center")
    c2 = wish.cell(4, col, f)
    c2.font = font(14, True, C["wish"])
    c2.number_format = fmt
    c2.alignment = Alignment(horizontal="center")
    c2.fill = fill("FFF0F5")
    c2.border = BORDER
wish.row_dimensions[4].height = 26

for i in range(WISH_ROWS):
    r = W_FIRST + i
    n = i + 1
    m = f"MATCH({n},exp_wishno,0)"
    wish.cell(r, 1, f'=IF(ISNA({m}),"",{n})')
    for c, nm in ((2, "exp_desc"), (3, "exp_cat"), (4, "exp_prio"), (5, "exp_year"),
                  (6, "exp_month"), (7, "exp_plan"), (8, "exp_fact"), (9, "exp_status")):
        wish.cell(r, c, f'=IF($A{r}="","",IF(INDEX({nm},{m})="","",INDEX({nm},{m})))')
    for c in range(1, 10):
        cell = wish.cell(r, c)
        cell.border = BORDER
        cell.font = font()
        cell.alignment = Alignment(vertical="center", horizontal="center" if c in (1, 4, 5, 6) else None)
    wish.cell(r, 7).number_format = MONEY
    wish.cell(r, 8).number_format = MONEY
    wish.cell(r, 7).fill = fill(C["plan_fill"])
    wish.cell(r, 8).fill = fill(C["fact_fill"])
wr = f"A{W_FIRST}:I{W_LAST}"
wish.conditional_formatting.add(wr, FormulaRule(
    formula=[f'$I{W_FIRST}="{DONE_EXP}"'], stopIfTrue=True,
    font=Font(strike=True, color="8D99AE"), fill=fill(C["done_fill"])))
wish.conditional_formatting.add(wr, FormulaRule(
    formula=[f'$I{W_FIRST}="{CANCEL_EXP}"'], stopIfTrue=True,
    font=Font(strike=True, italic=True, color="C0C0C0")))
wish.conditional_formatting.add(wr, FormulaRule(
    formula=[f'$C{W_FIRST}="{CATS[4][0]}"'], fill=fill("FFF4D6")))
wish.conditional_formatting.add(wr, FormulaRule(
    formula=[f'$C{W_FIRST}="{CATS[3][0]}"'], fill=fill("F3E8FF")))
wish.conditional_formatting.add(wr, FormulaRule(
    formula=[f'$C{W_FIRST}="{CATS[2][0]}"'], fill=fill("FFEAF1")))
wish.conditional_formatting.add("D4", DataBarRule(
    start_type="num", start_value=0, end_type="num", end_value=1, color="FF5D8F"))
wish.freeze_panes = "C7"

# ---------------------------------------------------------------- ГЛАВНАЯ
title(home, "🏠 Финансовый план 2026 – 2030",
      "Личный планер доходов и расходов: план vs факт по месяцам и годам, хотелки и мечты.",
      C["home"], 12)
widths(home, [3, 26, 18, 18, 18, 3, 26, 18, 18, 18, 3, 3])

def block_title(row, col, text, color):
    cell = home.cell(row, col, text)
    cell.font = Font(name=FONT, size=12, bold=True, color=color)


# KPI: весь период
block_title(4, 2, "📌 Итого за 2026–2030", C["home"])
for j, h in enumerate(["", "📝 ПЛАН", "✅ ФАКТ", "% выполнения"]):
    cell = home.cell(5, 2 + j, h)
    cell.font = Font(name=FONT, bold=True, color="FFFFFF")
    cell.fill = fill([C["home"], C["plan_head"], C["fact_head"], C["home"]][j])
    cell.alignment = Alignment(horizontal="center")
YR = f"{q(S_YEAR)}!"
kpi = [("💰 Доходы", f"={YR}B{YT}", f"={YR}C{YT}"),
       ("💸 Расходы", f"={YR}E{YT}", f"={YR}F{YT}"),
       ("⚖️ Остаток", f"={YR}H{YT}", f"={YR}I{YT}")]
for i, (lbl, p, f) in enumerate(kpi):
    r = 6 + i
    home.cell(r, 2, lbl).font = font(11, True)
    home.cell(r, 3, p)
    home.cell(r, 4, f)
    home.cell(r, 5, f'=IF(C{r}=0,"",D{r}/C{r})')
    home.cell(r, 3).fill = fill(C["plan_fill"])
    home.cell(r, 4).fill = fill(C["fact_fill"])
    for c in (3, 4):
        home.cell(r, c).number_format = MONEY
        home.cell(r, c).font = Font(name=FONT, size=11, bold=True, color="008000")
    home.cell(r, 5).number_format = PCT
    for c in range(2, 6):
        home.cell(r, c).border = BORDER
    home.row_dimensions[r].height = 22

# KPI: текущий месяц
block_title(4, 7, "🗓 Текущий месяц", C["mon"])
home.cell(4, 9, '=INDEX(lst_months,MONTH(TODAY()))&" "&YEAR(TODAY())').font = font(11, True, C["mon"])
for j, h in enumerate(["", "📝 ПЛАН", "✅ ФАКТ", "% выполнения"]):
    cell = home.cell(5, 7 + j, h)
    cell.font = Font(name=FONT, bold=True, color="FFFFFF")
    cell.fill = fill([C["mon"], C["plan_head"], C["fact_head"], C["mon"]][j])
    cell.alignment = Alignment(horizontal="center")
MN = f"{q(S_MON)}!"
# позиция текущего месяца в листе «По месяцам»
pos = f"((YEAR(TODAY())-{YEARS[0]})*12+MONTH(TODAY()))"
mkpi = [("💰 Доходы", "D", "E"), ("💸 Расходы", "G", "H"), ("⚖️ Остаток", "J", "K")]
for i, (lbl, pc, fc) in enumerate(mkpi):
    r = 6 + i
    home.cell(r, 7, lbl).font = font(11, True)
    home.cell(r, 8, f'=IFERROR(INDEX({MN}${pc}${M_FIRST}:${pc}${M_LAST},{pos}),0)')
    home.cell(r, 9, f'=IFERROR(INDEX({MN}${fc}${M_FIRST}:${fc}${M_LAST},{pos}),0)')
    home.cell(r, 10, f'=IF(H{r}=0,"",I{r}/H{r})')
    home.cell(r, 8).fill = fill(C["plan_fill"])
    home.cell(r, 9).fill = fill(C["fact_fill"])
    for c in (8, 9):
        home.cell(r, c).number_format = MONEY
        home.cell(r, c).font = Font(name=FONT, size=11, bold=True, color="008000")
    home.cell(r, 10).number_format = PCT
    for c in range(7, 11):
        home.cell(r, c).border = BORDER
for rng_ in ("E6:E8", "J6:J8"):
    home.conditional_formatting.add(rng_, DataBarRule(
        start_type="num", start_value=0, end_type="num", end_value=1, color="4B3F9E"))
home.conditional_formatting.add("C8:D8", CellIsRule(operator="lessThan", formula=["0"],
                                                    font=Font(bold=True, color=C["neg"])))
home.conditional_formatting.add("H8:I8", CellIsRule(operator="lessThan", formula=["0"],
                                                    font=Font(bold=True, color=C["neg"])))

# Хотелки
block_title(10, 2, "✨ Хотелки", C["wish"])
WS = f"{q(S_WISH)}!"
for j, (lbl, ref_cell, fmt) in enumerate([("Всего", "B4", "0"), ("Исполнено", "C4", "0"),
                                          ("Прогресс", "D4", PCT)]):
    home.cell(11, 3 + j, lbl).font = font(9, True, C["muted"])
    home.cell(11, 3 + j).alignment = Alignment(horizontal="center")
    c = home.cell(12, 3 + j, f"={WS}{ref_cell}")
    c.number_format = fmt
    c.font = Font(name=FONT, size=14, bold=True, color=C["wish"])
    c.alignment = Alignment(horizontal="center")
    c.fill = fill("FFF0F5")
    c.border = BORDER
home.conditional_formatting.add("E12", DataBarRule(
    start_type="num", start_value=0, end_type="num", end_value=1, color="FF5D8F"))

# Навигация
block_title(10, 7, "🧭 Навигация", C["home"])
nav = [(S_EXP, "вносить траты (план и факт)"), (S_INC, "вносить доходы (план и факт)"),
       (S_MON, "сводка по месяцам + графики"), (S_YEAR, "сводка по годам и категориям"),
       (S_WISH, "список хотелок и мечт"), (S_REF, "справочник для выпадающих списков")]
for i, (sh, desc) in enumerate(nav):
    c = home.cell(11 + i, 7, sh)
    c.hyperlink = f"#'{sh}'!A1"
    c.font = Font(name=FONT, size=11, bold=True, color="1155CC", underline="single")
    home.cell(11 + i, 8, desc).font = font(9, italic=True, color=C["muted"])

# Легенда
block_title(18, 2, "🎨 Как пользоваться", C["home"])
legend = [
    (C["plan_fill"], "📝 ПЛАН — синие колонки. Заполняете заранее: сколько собираетесь потратить / потенциально заработать."),
    (C["fact_fill"], "✅ ФАКТ — зелёные колонки. Заполняете по факту: сколько реально потратили / получили."),
    (C["input_fill"], "Жёлтые ячейки в «Справочнике» — ваши статьи, категории и типы дохода. Синий шрифт = ввод вручную."),
    (C["done_fill"], "Статус «✅ Выполнено» / «✅ Получено» автоматически зачёркивает строку (и хотелку в списке)."),
    (C["today"], "Жёлтая строка в «По месяцам» — текущий месяц. Красный шрифт — перерасход или минус."),
    ("FFFFFF", "Листы «По месяцам», «По годам», «Хотелки» и «Главная» считаются сами — в них ничего вводить не нужно."),
    ("FFFFFF", "В листах «Расходы» и «Доходы» есть строки-ПРИМЕРЫ — замените или удалите их."),
]
for i, (bg, text) in enumerate(legend):
    r = 19 + i
    sw = home.cell(r, 2)
    sw.fill = fill(bg)
    sw.border = BORDER
    home.merge_cells(start_row=r, start_column=3, end_row=r, end_column=10)
    t = home.cell(r, 3, text)
    t.font = font(10)
    t.alignment = Alignment(vertical="center", wrap_text=True)
    home.row_dimensions[r].height = 20
home.cell(19, 2, "ПЛАН").font = font(9, True, C["plan_head"])
home.cell(20, 2, "ФАКТ").font = font(9, True, C["fact_head"])
home.cell(21, 2, "ввод").font = font(9, True, "0000FF")
home.cell(22, 2, "зачёркнуто").font = Font(name=FONT, size=9, strike=True, color="8D99AE")
home.cell(23, 2, "сейчас").font = font(9, True)
for r in range(19, 24):
    home.cell(r, 2).alignment = Alignment(horizontal="center", vertical="center")

# Графики на главной
ch = line_chart("💰 Доход по месяцам: план vs факт", "₸", height=8, width=17)
ch.add_data(Reference(mon, min_col=4, max_col=5, min_row=4, max_row=M_LAST), titles_from_data=True)
ch.set_categories(cats_ref)
color_series(ch, [C["plan_head"], C["fact_head"]], dashed=[0])
home.add_chart(ch, "B27")
ch = line_chart("📈 Доход и расход по годам (план/факт)", "₸", height=8, width=17)
ch.add_data(Reference(yr, min_col=2, max_col=3, min_row=4, max_row=Y_LAST), titles_from_data=True)
ch.add_data(Reference(yr, min_col=5, max_col=6, min_row=4, max_row=Y_LAST), titles_from_data=True)
ch.set_categories(ych_cats)
color_series(ch, [C["plan_head"], C["fact_head"], "F4A261", C["neg"]], dashed=[0, 2])
home.add_chart(ch, "G27")

wb.active = 0
wb.save(OUT)
print(f"Saved {OUT}")
